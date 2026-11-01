"""Deterministic, source-bound offline reports. Model prose is data, never a template."""
from __future__ import annotations

import csv
import io
import json
import math
import re
import shutil
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

from bs4 import BeautifulSoup
from jinja2 import Environment, FileSystemLoader, StrictUndefined, select_autoescape
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Font, PatternFill
from PIL import Image
from src.chart_svg import CHART_RENDERING, chart_csp, inline_svg, interaction_script

from src.answers import identifier_tokens, unbound_numbers
from src.text_guard import CN_NUMBER, NUMERIC_LEXEME, TOKEN


def _normalized_numbers(text: str):
    """Numeric values with formatting noise removed (thousand separators,
    trailing decimal zeros) so formatting-only edits pass the gate."""
    from decimal import Decimal, InvalidOperation
    values = []
    for lex in NUMERIC_LEXEME.findall(text):
        clean = lex.replace(',', '')
        try:
            values.append(Decimal(clean).normalize())
        except InvalidOperation:
            values.append(lex)
    return values

from src.schemas import (AnalysisResult, ReportChartEdit, ReportSectionEdit, ReportLayout, ReportRevisionPatch,
                         ReportResult, ReportTextReplacement, TaskSpec, new_id)
from pydantic import ValidationError
from src.storage import Store, StoreError, file_hash
from src.performance import span, traced_method
from src.validation import validation_scope, validated_method

REPORT_STYLE_REVISION = 'responsive-20261024'
XML_BAD = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')
# A bounded static guard against accidental unbound numerical claims, not a
# proof that free-form qualitative prose follows logically from the data.
NUMBER = re.compile(r'\d|百分之|[' + CN_NUMBER + r']{2,}|负[' + CN_NUMBER + r']|[' + CN_NUMBER + r']+点[' + CN_NUMBER + r']+|[' + CN_NUMBER + r'](?:倍|成|个|项|行|列|分|元|[%％])|'
    r'(?:为|是|等于|约|达到|超过|低于|高于|至少|最多|合计|共|均值|平均值|中位数|总量|总额|数量|比例|占比|增加|减少|下降|提升)\s*[' +
    CN_NUMBER + r'](?=$|[。；，、%％\s）)])')
HTML_TABLE_LIMIT = 200
VERSION = '1.0'


def exact(value) -> str:
    if isinstance(value, float) and not math.isfinite(value):
        raise StoreError('Non-finite report metric is not supported', 422, 'invalid_metric')
    if isinstance(value, (dict, list, bool)) or value is None:
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    return str(value)


class NeedAnalysis(Exception):
    def __init__(self, message: str, reason_code: str):
        super().__init__(message)
        self.reason_code = reason_code


class ReportService:
    def __init__(self, store: Store):
        self.store = store
        self.cleaned_characters = 0
        self._database_context = {}
        # Official Jinja Environment/select_autoescape pattern; only the
        # application-owned template is loaded, never model-supplied source.
        # https://jinja.palletsprojects.com/en/stable/api/#autoescaping
        self.environment = Environment(loader=FileSystemLoader(Path(__file__).resolve().parents[1] / 'templates'),
            autoescape=select_autoescape(['html', 'xml']), undefined=StrictUndefined)

    @traced_method('report_read')
    @validated_method
    def read_report(self, task: TaskSpec, report_id: str) -> dict:
        """Load a registered successful same-session version, never arbitrary HTML."""
        artifact = self.store.get_artifact(report_id)
        if artifact.session_id != task.session_id or artifact.kind != 'report':
            raise StoreError('Report is not an authorized generated report in this session', 404, 'not_found')
        source_task = self.store.get_task(artifact.task_id)
        if source_task['status'] != 'succeeded' or report_id not in (source_task.get('result') or {}).get('report_ids', []):
            raise StoreError('Only successfully delivered reports may be revised', 422, 'invalid_report_base')
        path = self.store.resolve_artifact(report_id)
        candidates = [item for item in self.store.list_artifacts(artifact.task_id)
                      if item.kind == 'report_manifest' and Path(item.relative_path).parent == Path(artifact.relative_path).parent]
        if len(candidates) != 1:
            raise StoreError('Report version has no unique registered manifest', 422, 'invalid_report_base')
        manifest = json.loads(self.store.resolve_artifact(candidates[0].artifact_id).read_text(encoding='utf-8'))
        if (manifest.get('task_id') != artifact.task_id or manifest.get('session_id') != task.session_id or
                not any(item['name'] == path.name and item['sha256'] == artifact.hash for item in manifest.get('files', []))):
            raise StoreError('Report manifest provenance mismatch', 422, 'invalid_report_base')
        ids = [item['analysis_id'] for item in manifest['analyses']]
        current = self._analyses(task, ids)
        try:
            # Default-only schema additions (such as db_refs=[]) must not make
            # an unchanged historical file report stale. Validate every field;
            # unknown fields and altered metrics/provenance still fail closed.
            snapshots = [AnalysisResult.model_validate(item).model_dump() for item in manifest['analyses']]
        except ValidationError as exc:
            raise StoreError('Base report analysis snapshot is invalid', 422, 'invalid_report_base') from exc
        if [item.model_dump() for item in current] != snapshots:
            raise StoreError('Base report analysis snapshot has changed', 409, 'stale_report_analysis')
        presentation = manifest.get('presentation')
        if presentation is None:
            # Backward compatibility for versions created before presentation
            # snapshots existed. Read only hash-checked generated files.
            siblings = [item for item in self.store.list_artifacts(artifact.task_id)
                        if item.kind == 'report' and Path(item.relative_path).parent == Path(artifact.relative_path).parent
                        and Path(item.relative_path).suffix == '.html']
            if siblings:
                soup = BeautifulSoup(self.store.resolve_artifact(siblings[0].artifact_id).read_text(encoding='utf-8'), 'html.parser')
                heading = soup.find('h2', string='结论') or soup.find('h2', string='摘要')
                summary = heading.find_next_sibling('p').get_text() if heading else ''
                sections = []
                for article in soup.find_all('article'):
                    refs = []
                    for row in article.select('table tr')[1:]:
                        cells = row.find_all('td')
                        if len(cells) == 3:
                            aid, _, key = cells[2].get_text().partition(':')
                            if not any(a['analysis_id'] == aid and key in a['metrics'] for a in manifest['analyses']):
                                raise StoreError('Legacy chapter metric reference is invalid', 422, 'invalid_report_base')
                            refs.append(dict(analysis_id=aid, key=key, label=cells[0].get_text(), value=cells[1].get_text()))
                    sections.append(dict(heading=article.h2.get_text(), text=article.p.get_text(), metrics=refs,
                                         table_ids=[], chart_ids=[]))
                presentation = dict(title=soup.h1.get_text(), summary=summary, sections=sections)
                legacy_metrics = soup.find('h2', string='指标汇总') is not None
                legacy_tables = soup.find('h2', string='结果明细') is not None
            else:
                book = load_workbook(path, read_only=True, data_only=False)
                try:
                    info = dict(list(book['说明与口径'].values)[1:])
                    sections = []
                    for name in book.sheetnames:
                        if not name.startswith('章节_'):
                            continue
                        rows = list(book[name].values)[1:]
                        refs = []
                        for label, value in rows[2:]:
                            matches = [m for m in manifest['metric_bindings'] if m['value'] == str(value)]
                            if not matches:
                                raise StoreError('Legacy chapter metric reference is unavailable', 422, 'invalid_report_base')
                            refs.append(dict(analysis_id=matches[0]['analysis_id'], key=matches[0]['key'], label=label, value=str(value)))
                        sections.append(dict(heading=rows[0][1], text=rows[1][1] or '', metrics=refs, table_ids=[], chart_ids=[]))
                    presentation = dict(title=info['标题'], summary=info.get('结论', info.get('摘要', '')), sections=sections)
                    legacy_metrics = '指标汇总' in book.sheetnames
                    legacy_tables = any(name.startswith('结果表_') for name in book.sheetnames)
                finally:
                    book.close()
            presentation.update(layout=ReportLayout(include_metrics_summary=legacy_metrics,
                include_tables=legacy_tables).model_dump(), limitations=manifest.get('limitations', []))
        try:
            # Older snapshots contain only the original four switches. Add
            # defaults without restoring any previously removed chapters.
            presentation = {**presentation, 'layout': ReportLayout.model_validate({
                'include_tables': True, **presentation.get('layout', {})}).model_dump()}
        except ValidationError as exc:
            raise StoreError('Base report layout is invalid', 422, 'invalid_report_base') from exc
        # Older manifests did not snapshot these visible source-derived texts.
        # New versions persist them so a later edit never restores the originals.
        if 'findings' not in presentation:
            presentation['findings'] = list(dict.fromkeys(self._clean(x) for a in current
                if not any(ref.get('source_type') == 'material' for ref in a.kb_refs) for x in a.findings))
        if 'quality' not in presentation:
            presentation['quality'] = self._source_quality(task, current)
        presentation.setdefault('limitations_fallback', True)
        from src.report_edits import validate_override
        overrides = manifest.get('chart_overrides', [])
        if not isinstance(overrides, list) or len({item.get('chart_id') for item in overrides if isinstance(item, dict)}) != len(overrides):
            raise StoreError('局部图表来源记录无效。', 422, 'invalid_report_base')
        for override in overrides:
            validate_override(self.store, task, current, override)
        from src.report_revision import initial_state, validate_state
        component_state = validate_state(self.store, task, current,
            manifest.get('component_state') or initial_state(self.store, current, presentation['layout'], overrides), overrides)
        by_chart = {item['chart_id']: item for item in overrides}
        charts = []
        for slot in component_state['charts']:
            cid, current_id = slot['component_id'], slot['artifact_id']
            analysis = next(a for a in current if a.analysis_id == slot['analysis_id'])
            override = by_chart.get(cid)
            if override and override['artifact_id'] != current_id:
                override = None
            chart_artifact = self.store.get_artifact(current_id)
            chart_path = self.store.resolve_artifact(current_id)
            chart_title = Path(chart_artifact.name).stem
            if chart_path.suffix.lower() == '.svg':
                from src.chart_svg import parse_svg
                chart_title = parse_svg(chart_path.read_bytes()).get('data-chart-title') or chart_title
            charts.append(dict(chart_id=cid, current_artifact_id=current_id, title=chart_title,
                name=chart_artifact.name, analysis_id=analysis.analysis_id,
                edit_spec=override['spec'] if override else None,
                table_ids=analysis.table_ids, visible=slot['visible']))
        return dict(report_id=report_id, task_id=artifact.task_id, version=manifest['report_version'], hash=artifact.hash,
                    analysis_ids=ids, formats=manifest['formats'], presentation=presentation,
                    charts=charts, chart_overrides=overrides, component_state=component_state,
                    chart_rendering=manifest.get('chart_rendering'),
                    needs_visual_upgrade=bool('html' in manifest['formats'] and (
                        manifest.get('style_revision') != REPORT_STYLE_REVISION or
                        (any(a.chart_ids for a in current) and manifest.get('chart_rendering') != CHART_RENDERING))))

    def _clean(self, value) -> str:
        text = str(value)
        clean, removed = XML_BAD.subn('', text)
        self.cleaned_characters += removed
        if len(clean) > 32767:
            raise StoreError('Report text exceeds Excel cell length; split the source text into sections', 422, 'report_text_too_long')
        return clean

    def _analyses(self, task, ids):
        record = self.store.get_task(task.task_id)
        if record['session_id'] != task.session_id:
            raise StoreError('Report task is not in this session', 404, 'not_found')
        self.store.get_session(task.session_id)
        allowed = set(task.analysis_ids)
        for prior in self.store.list_tasks():
            if prior['session_id'] != task.session_id:
                continue
            if prior['status'] == 'succeeded':
                allowed.update((prior.get('result') or {}).get('analysis_ids', []))
        result = []
        for aid in dict.fromkeys(ids):
            analysis = self.store.get_analysis(aid, task.session_id)
            if analysis.session_id != task.session_id or self.store.get_task(analysis.task_id)['session_id'] != task.session_id:
                raise StoreError('Analysis provenance is not in this session', 404, 'not_found')
            if aid not in allowed and analysis.task_id != task.task_id:
                raise StoreError('Analysis is not authorized for this report task', 422, 'analysis_not_authorized')
            if analysis.status != 'completed' or not (analysis.input_versions or analysis.db_refs):
                raise StoreError('Report requires completed analysis with source versions', 422, 'invalid_analysis')
            if task.file_ids and not set(analysis.input_versions).issubset(task.file_ids):
                raise StoreError('Analysis sources exceed the selected task files', 422, 'file_not_selected')
            for artifact_id, kind in [(x, 'table') for x in analysis.table_ids] + [(x, 'chart') for x in analysis.chart_ids] + [(x, 'script') for x in analysis.script_ids]:
                artifact = self.store.get_artifact(artifact_id)
                if artifact.session_id != task.session_id or artifact.task_id != analysis.task_id or artifact.kind != kind:
                    raise StoreError('Analysis artifact provenance mismatch', 422, 'invalid_artifact')
                self.store.resolve_artifact(artifact_id)
            # Store validates the originating task. A report must additionally
            # retain authorization under its own current source/table selection.
            self._database_snapshots(task, analysis, refresh=True)
            result.append(analysis)
        return result

    def _database_snapshots(self, task, analysis, *, refresh=False):
        if not analysis.db_refs:
            return []  # File-only reports remain available when PostgreSQL is down.
        key = (task.task_id, analysis.analysis_id,
               json.dumps(analysis.db_refs, sort_keys=True, separators=(',', ':')))
        if refresh or key not in self._database_context:
            from src.database import DatabaseService
            service = DatabaseService(self.store)
            self._database_context[key] = [service.authorize_snapshot(task, ref) for ref in analysis.db_refs]
        return self._database_context[key]

    def _source_quality(self, task, analyses):
        quality = []
        seen = set()
        for analysis in analyses:
            for fid in analysis.input_versions:
                if fid in seen:
                    continue
                seen.add(fid)
                file = self.store.get_file(fid, task.session_id)
                profiles = file.get('profiles') or ([file['profile']] if file.get('profile') else [])
                for profile in profiles:
                    for warning in profile.get('quality_warnings', []):
                        item = dict(file_id=fid, table=self._clean(profile.get('sheet_or_table') or ''), text=self._clean(warning))
                        if item not in quality:
                            quality.append(item)
                if not profiles:
                    quality.append(dict(file_id=fid, table='', text='未登记源文件质量概览；缺失处理请见分析口径，报告未另行推断数据质量。'))
            for snapshot in self._database_snapshots(task, analysis):
                key = ('postgresql', snapshot['query_id'])
                if key in seen:
                    continue
                seen.add(key)
                quality.append(dict(source_type='postgresql', source_id=snapshot['source_id'],
                    query_id=snapshot['query_id'], table=', '.join(snapshot['tables']),
                    text=self._clean(f"使用已登记的完整 SQL 结果快照；查询时间 {snapshot['queried_at']}；"
                        f"权限范围：{snapshot['scope_description']}。未另行扫描数据库或推断缺失、重复等质量指标。")))
        return quality

    def _replace_text(self, context, replacements):
        """Validate every match before mutating, then replace in place, never append."""
        fields = {key: [(context, key)] for key in ('title', 'summary')}
        for key in ('findings', 'limitations'):
            fields[key] = [(context[key], i) for i in range(len(context[key]))]
        fields['quality'] = [(item, 'text') for item in context['quality']]
        fields['section_heading'] = [(item, 'heading') for item in context['sections']]
        fields['section_text'] = [(item, 'text') for item in context['sections']]
        fields['metric_label'] = [(metric, 'label') for section in context['sections'] for metric in section['metrics']]
        pending, seen, applied = [], {}, []
        for value in replacements:
            try:
                item = ReportTextReplacement.model_validate(value)
            except ValidationError as exc:
                raise StoreError('翻译替换参数无效。', 422, 'invalid_text_replacement') from exc
            identity = (item.field, item.original)
            if identity in seen:
                if seen[identity] != item.replacement:
                    raise StoreError('同一原文不能指定互相冲突的译文。', 422, 'invalid_text_replacement')
                continue  # Shared quality notes can occur in several files.
            seen[identity] = item.replacement
            matches = [(container, key) for container, key in fields[item.field] if container[key] == item.original]
            if not matches:
                raise StoreError('原文未匹配所选报告，请重新 read_report 并使用完整原文。', 422, 'text_not_found')
            replacement = self._clean(item.replacement)
            if not replacement.strip():
                raise StoreError('译文不能为空。', 422, 'invalid_text_replacement')
            if Counter(_normalized_numbers(item.original)) != Counter(_normalized_numbers(replacement)):
                raise StoreError('翻译不能增加、删除或改变原文数值；请保留原始数字与精度。', 422, 'translation_number_changed')
            pending.extend((container, key, replacement) for container, key in matches)
            if replacement != item.original:
                applied.append(item.model_dump())
        for container, key, replacement in pending:
            container[key] = replacement
        for field in ('findings', 'limitations'):
            context[field] = list(dict.fromkeys(context[field]))
        context['quality'] = [dict(values) for values in dict.fromkeys(tuple(item.items()) for item in context['quality'])]
        return applied

    def _metric(self, aid, key, analyses):
        if aid not in analyses or key not in analyses[aid].metrics:
            raise NeedAnalysis(f'缺少已验证指标：{aid}:{key}，请补充分析后再次生成报告。', 'missing_metric')
        return self._clean(exact(analyses[aid].metrics[key]))

    def _prose(self, value, analyses):
        if not isinstance(value, str):
            raise StoreError('Report prose must be text', 422, 'invalid_narrative')
        residue = TOKEN.sub('', value)
        if '{{metric:' in residue:
            raise NeedAnalysis('指标占位符格式错误，请使用 {{metric:analysis_id:metric_key}}。', 'invalid_placeholder')
        offenders = unbound_numbers(residue, identifier_tokens(analyses.values()))
        if offenders:
            raise NeedAnalysis('报告文字包含未绑定数值：' + '、'.join(offenders[:6]) +
                '；统计数值请改用 metric_refs 或指标占位符引用已验证指标，'
                '数据中真实存在的年份、分组名、列名等标识可直接书写。', 'unbound_prose')
        return self._clean(TOKEN.sub(lambda m: self._metric(m[1], m[2], analyses), value))

    def _default_summary(self, analyses, request=''):
        """State registered results without deriving statistics in the renderer.

        This path also serves the safe narrative fallback. Table previews and raw
        rows are never used to infer ranks, proportions, trends or causality.
        """
        records = self._conclusion_metrics(analyses)
        comparison = self._comparison_conclusion(records, request)
        if comparison:
            return comparison
        insufficient = bool(re.search(r'比较|对比|差异|关系|高低|compare|versus', request, re.I)
            and not any(re.search(r'corr|pearson|spearman|difference|slope', item['key'], re.I) for item in records))
        boundary = '现有已登记指标不足以完成所问的比较或关系判断；以下仅列可核实事实。\n' if insufficient else ''
        facts = []
        # Findings can contain rounded ratios or an inferred currency even when
        # their source metrics are sound. Only promote exact registered values;
        # never turn free-form numerical claims into a fallback escape hatch.
        values = {exact(value) for analysis in analyses for value in analysis.metrics.values()
                  if value is not None and not isinstance(value, (dict, list, bool))}
        for analysis in analyses:
            if any(ref.get('source_type') == 'material' for ref in analysis.kb_refs):
                continue
            for finding in analysis.findings:
                text = self._clean(finding).strip()
                # Fixed analysis tools record execution metadata as findings;
                # that is useful provenance but is not a business conclusion.
                if not text or re.search(r'source rows.*output rows|full result stored|'
                        r'已完成.*(?:计算|分析)|已验证分析结果整理|结果已保存|分析已完成', text, re.I):
                    continue
                if '{{metric:' in text:
                    text = self._prose(text, {a.analysis_id: a for a in analyses})
                elif any(number not in values for number in NUMERIC_LEXEME.findall(text)):
                    continue
                if NUMBER.search(NUMERIC_LEXEME.sub('', text)):
                    continue  # Chinese number words must not bypass metric binding.
                if re.search(r'\b(?:BRL|USD|EUR|CNY|RMB)\b|[$€¥￥]|人民币|美元|欧元|雷亚尔', text, re.I):
                    continue  # A finding is not an authoritative unit registry.
                if any(re.fullmatch(re.escape(key) + r'\s*(?:为|=|:)\s*' + re.escape(exact(value)) + r'[。.]?', text)
                       for key, value in analysis.metrics.items()):
                    continue
                if text not in facts:
                    facts.append(text)
        if facts:
            return boundary + '\n'.join(facts[:3])
        records.sort(key=lambda item: self._conclusion_relevance(item, request), reverse=True)
        for item in records[:3]:
            facts.append(self._clean(f"{item['label']}为{exact(item['value'])}。"))
        if facts:
            return boundary + '\n'.join(facts)
        if any(any(ref.get('source_type') == 'material' for ref in a.kb_refs) for a in analyses):
            return '本报告整理所选文档的原文内容；文档中的业务结论与数值尚未经过独立统计验证。'
        # A row listing is not evidence for a business answer. Explicit detail
        # exports still receive an honest scope statement, never a fabricated
        # ranking or inference from those rows.
        return '现有结果仅包含记录明细或处理范围，尚不足以形成业务结论；需补充直接回答问题的分析指标。'

    def _conclusion_metrics(self, analyses):
        """Select source-bound scalar values; labels contribute no new numbers."""
        metadata = {'source_rows', 'result_rows', 'right_rows', 'source_characters', 'excerpt_characters',
                    'left_duplicate_keys', 'right_duplicate_keys', 'left_key_match_rate', 'row_expansion_ratio'}
        words = {'avg': '平均', 'mean': '平均', 'median': '中位数', 'payment': '支付金额',
                 'revenue': '销售金额', 'sales': '销售额', 'latency': '响应时间', 'cost': '成本',
                 'repeat': '复购客户', 'single': '单次客户', 'total': '总', 'customers': '客户数',
                 'orders': '订单数', 'count': '数量', 'sum': '合计', 'min': '最小值', 'max': '最大值'}
        records = []
        seen = set()
        for analysis in analyses:
            if any(ref.get('source_type') == 'material' for ref in analysis.kb_refs):
                continue
            for key, value in analysis.metrics.items():
                if key in metadata or value is None or isinstance(value, (dict, list, bool)):
                    continue
                definition = analysis.metric_definitions.get(key, {})
                label = definition.get('label') if isinstance(definition, dict) else None
                tokens = key.lower().split('_')
                if (not isinstance(label, str) or not label.strip() or NUMBER.search(label)
                        or re.search(r'\b(?:BRL|USD|EUR|CNY|RMB)\b|[$€¥￥]', label, re.I)):
                    stem, _, suffix = key.rpartition('_')
                    suffixes = {'mean': '平均值', 'sum': '合计', 'median': '中位数', 'min': '最小值',
                                'max': '最大值', 'std': '标准差', 'count': '有效记录数', 'nunique': '去重数量'}
                    label = (f'{stem} {suffixes[suffix]}' if stem and suffix in suffixes else
                             ''.join(words.get(token, token) for token in tokens))
                if NUMBER.search(label):
                    continue
                prefix = ''
                if len(analyses) > 1 and len(analysis.input_versions) == 1:
                    fid = next(iter(analysis.input_versions))
                    prefix = Path(self.store.get_file(fid, analysis.session_id)['name']).stem + '：'
                    if NUMBER.search(prefix):
                        prefix = ''
                identity = (key, prefix, tuple(sorted({ref.get('source_id', '') for ref in analysis.db_refs})))
                if identity in seen:
                    continue  # Preserve the first registered, full-precision result.
                seen.add(identity)
                records.append(dict(key=key, label=self._clean(prefix + label), tokens=tokens, value=value,
                                    unit=definition.get('unit') if isinstance(definition, dict) else None))
        return records

    @staticmethod
    def _conclusion_relevance(item, request):
        aliases = {'avg': ('平均', '人均', '均值'), 'mean': ('平均', '人均', '均值'),
                   'payment': ('支付', '付款'), 'revenue': ('销售', '收入'), 'sales': ('销售',),
                   'latency': ('响应时间', '耗时', '延迟'), 'cost': ('成本',), 'count': ('数量', '多少'),
                   'median': ('中位数',), 'repeat': ('复购',), 'single': ('单次',)}
        request = request.casefold()
        score = sum(3 for token in item['tokens'] for term in (*aliases.get(token, ()), token) if term in request)
        label = item['label']
        score += sum(1 for index in range(len(label) - 1)
                     if re.fullmatch(r'[\u4e00-\u9fff]{2}', label[index:index + 2]) and label[index:index + 2] in request)
        return score

    def _comparison_conclusion(self, records, request):
        if not re.search(r'比较|对比|差异|关系|高低|相比|compare|versus|\bvs\b', request, re.I):
            return None
        candidates = []
        for index, left in enumerate(records):
            for right in records[index + 1:]:
                # Different cohorts, same registered measure. Do not compare a
                # count with an average or derive statistics from table rows.
                a, b = left['tokens'], right['tokens']
                suffix = 0
                while suffix < min(len(a), len(b)) and a[-suffix - 1] == b[-suffix - 1]:
                    suffix += 1
                if suffix == len(a) == len(b) and left['label'] != right['label']:
                    measure = a  # Same statistic, separately identified input sources.
                elif suffix and suffix < len(a) and suffix < len(b):
                    measure = a[-suffix:]
                else:
                    prefix = 0
                    while prefix < min(len(a), len(b)) and a[prefix] == b[prefix]:
                        prefix += 1
                    if not prefix or prefix == len(a) or prefix == len(b):
                        continue
                    measure = a[:prefix]
                if not any(word in measure for word in ('avg', 'mean', 'median', 'payment', 'revenue', 'sales', 'latency', 'cost')):
                    continue
                operators = {'avg': 'mean', 'mean': 'mean', 'sum': 'sum', 'total': 'sum',
                             'median': 'median', 'count': 'count', 'min': 'min', 'max': 'max'}
                if not any(token not in operators for token in measure):
                    continue  # Shared AVG alone does not establish the same quantity.
                if [operators[token] for token in a if token in operators] != [operators[token] for token in b if token in operators]:
                    continue
                if left['unit'] and right['unit'] and left['unit'] != right['unit']:
                    continue
                try:
                    av, bv = Decimal(str(left['value'])), Decimal(str(right['value']))
                except (InvalidOperation, ValueError):
                    continue
                if not av.is_finite() or not bv.is_finite():
                    continue
                score = self._conclusion_relevance(left, request) + self._conclusion_relevance(right, request)
                if score <= 0:
                    continue
                relation = '高于' if av > bv else '低于' if av < bv else '等于'
                text = f"{left['label']}为{exact(left['value'])}，{relation}{right['label']}的{exact(right['value'])}。"
                candidates.append((score, text))
        if not candidates:
            return None
        return self._clean(max(candidates, key=lambda item: item[0])[1] + '该比较描述已登记结果的差异，不代表因果关系。')

    def _table(self, artifact_id, aid, *, include_rows=True):
        artifact = self.store.get_artifact(artifact_id)
        path = self.store.resolve_artifact(artifact_id)
        if path.suffix.lower() != '.csv':
            raise StoreError('Result tables must be registered CSV artifacts', 422, 'invalid_table')
        rows = []
        count = 0
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.reader(stream, strict=True)
            columns = next(reader, None)
            if not columns:
                raise StoreError('Invalid result CSV shape', 422, 'invalid_table')
            for row in reader:
                if len(row) != len(columns):
                    raise StoreError('Invalid result CSV shape', 422, 'invalid_table')
                count += 1
                if include_rows:
                    rows.append(row)
        if include_rows and (len(columns) > 16384 or (count + 1) * len(columns) > 5000000):
            raise StoreError('Report table exceeds export cell limits; aggregate first', 422, 'report_too_large')
        # Keep every CSV lexeme: 001/NA/formula-like text must not be inferred as
        # missing data or formulas, and full tables must not use preview APIs.
        rows = [[self._clean(cell) for cell in row] for row in rows]
        return dict(artifact_id=artifact_id, analysis_id=aid, name=self._clean(artifact.name), hash=artifact.hash,
            columns=[self._clean(cell) for cell in columns], rows=rows, total_rows=count,
            shown_rows=min(count, HTML_TABLE_LIMIT) if include_rows else 0,
            download_url=f'/api/v1/artifacts/{artifact_id}/download')

    def _chart(self, artifact_id, aid):
        artifact = self.store.get_artifact(artifact_id)
        path = self.store.resolve_artifact(artifact_id)
        if path.suffix.lower() not in {'.png', '.svg'} or path.stat().st_size > 20 * 1024 * 1024:
            raise StoreError('Expected a registered SVG/PNG chart up to 20 MiB', 422, 'invalid_chart')
        pair_kind = 'chart_svg' if path.suffix.lower() == '.png' else 'chart_preview'
        pairs = [a for a in self.store.list_artifacts(artifact.task_id)
                 if a.kind == pair_kind and a.session_id == artifact.session_id and a.source_ids == [artifact_id]
                 and Path(a.relative_path) == path.relative_to(self.store.root).with_suffix('.svg' if pair_kind == 'chart_svg' else '.png')]
        if len(pairs) > 1:
            raise StoreError('Ambiguous chart companion', 422, 'invalid_chart')
        svg_artifact = artifact if path.suffix.lower() == '.svg' else (pairs[0] if pairs else None)
        png_artifact = artifact if path.suffix.lower() == '.png' else (pairs[0] if pairs else None)
        raw = None
        if png_artifact:
            raw = self.store.resolve_artifact(png_artifact.artifact_id).read_bytes()
            try:
                with Image.open(io.BytesIO(raw)) as picture:
                    if picture.format != 'PNG' or picture.width * picture.height > 40000000:
                        raise ValueError('Invalid chart dimensions or format')
                    picture.verify()
            except (OSError, ValueError, Image.DecompressionBombError) as exc:
                raise StoreError('Chart contents are not a valid PNG', 422, 'invalid_chart') from exc
        markup = None
        if getattr(self, 'require_svg', True):
            if not svg_artifact:
                raise NeedAnalysis('此图只有 PNG，无法无损变为矢量图。请仅为目标图补充有来源的 SVG 并在原版本上替换，原报告及未修改部分保留。', 'missing_vector_chart')
            try:
                markup = inline_svg(self.store.resolve_artifact(svg_artifact.artifact_id).read_bytes(), artifact_id)
            except (ValueError, OSError) as exc:
                raise StoreError('Invalid or unsafe vector chart: ' + str(exc), 422, 'invalid_chart') from exc
        if getattr(self, 'require_png', True) and raw is None:
            raise NeedAnalysis('Excel 图表需要与 SVG 同次导出的 PNG 预览。请用 Matplotlib 重新导出图表。', 'missing_chart_preview')
        return dict(artifact_id=artifact_id, analysis_id=aid, name=self._clean(artifact.name), hash=artifact.hash,
                    title=(ET.fromstring(str(markup)).get('data-chart-title') or
                           (Path(artifact.name).stem.replace('_', ' ').replace('-', ' ') if not re.fullmatch(r'[0-9a-f]{32}', Path(artifact.name).stem) else '分析图表')) if markup else '分析图表',
                    svg=markup, raw=raw, svg_artifact_id=svg_artifact.artifact_id if svg_artifact else None,
                    preview_artifact_id=png_artifact.artifact_id if png_artifact else None)

    def _context(self, task, analyses, title, narrative, version, component_state=None):
        indexed = {a.analysis_id: a for a in analyses}
        if narrative is not None and not isinstance(narrative, dict):
            raise StoreError('Narrative must be a structured object', 422, 'invalid_narrative')
        narrative = narrative or {}
        if set(narrative) - {'summary', 'sections', 'limitations'}:
            raise StoreError('Unknown narrative fields', 422, 'invalid_narrative')
        metrics, tables, charts, sources, materials = [], {}, {}, {}, []
        quality = self._source_quality(task, analyses)
        for analysis in analyses:
            material_refs = [ref for ref in analysis.kb_refs if ref.get('source_type') == 'material']
            for key, value in analysis.metrics.items():
                raw = self._clean(exact(value)) if isinstance(value, (str, dict, list)) else value
                metrics.append(dict(analysis_id=analysis.analysis_id, key=self._clean(key), label=self._clean(key),
                    value=self._metric(analysis.analysis_id, key, indexed), raw=raw,
                    definition=self._clean(json.dumps(analysis.metric_definitions.get(key, {}), ensure_ascii=False))))
            for ref in material_refs:
                if ref.get('file_id') not in analysis.input_versions or not isinstance(ref.get('text'), str):
                    raise StoreError('Material excerpt must reference an analysis source', 422, 'invalid_material')
                materials.append(dict(analysis_id=analysis.analysis_id, file_id=ref['file_id'],
                    location=self._clean(ref.get('location', 'extracted text')), text=self._clean(ref['text']),
                    truncated=bool(ref.get('truncated', False))))
            for fid, digest in analysis.input_versions.items():
                file = self.store.get_file(fid, task.session_id)
                self.store.resolve_file(fid, task.session_id)
                sources[fid] = dict(file_id=fid, name=self._clean(file['name']), hash=digest, source=file['source'])
            for snapshot in self._database_snapshots(task, analysis):
                metadata = self.store.get_artifact(snapshot['metadata_artifact_id'])
                sources[('postgresql', snapshot['query_id'])] = dict(source_type='postgresql', source='PostgreSQL',
                    source_id=snapshot['source_id'], query_id=snapshot['query_id'],
                    name=self._clean(f"{snapshot['database']} / {snapshot['source_id']}"),
                    hash=snapshot['result_hash'], result_hash=snapshot['result_hash'],
                    sql_hash=snapshot['sql_hash'], metadata_hash=metadata.hash,
                    queried_at=snapshot['queried_at'], snapshot_id=snapshot['snapshot_id'],
                    scope=self._clean(snapshot['scope_description']), scope_fingerprint=snapshot['scope_fingerprint'],
                    tables=snapshot['tables'], row_count=snapshot['row_count'],
                    sql=self._clean(snapshot['sql']), params=snapshot['params'],
                    **{key: snapshot[key] for key in ('result_artifact_id', 'sql_artifact_id', 'metadata_artifact_id')})
            for tid in analysis.table_ids:
                visible = component_state is None or any(s['artifact_id'] == tid and s['visible'] for s in component_state['tables'])
                checked = self._table(tid, analysis.analysis_id,
                    include_rows=getattr(self, 'include_result_tables', True) and visible)
                if visible:
                    tables[tid] = checked
            for cid in analysis.chart_ids:
                if component_state is None:
                    charts[cid] = self._chart(cid, analysis.analysis_id)
        if component_state is not None:
            tables = {s['artifact_id']: self._table(s['artifact_id'], s['analysis_id'],
                include_rows=getattr(self, 'include_result_tables', True))
                for s in component_state['tables'] if s['visible']}
            charts = {s['artifact_id']: self._chart(s['artifact_id'], s['analysis_id'])
                for s in component_state['charts'] if s['visible']}
        raw_sections = narrative.get('sections', [])
        if not isinstance(raw_sections, list) or len(raw_sections) > 100:
            raise StoreError('Report sections must be a list of at most 100 entries', 422, 'invalid_narrative')
        sections = []
        for section in raw_sections:
            if not isinstance(section, dict) or set(section) - {'heading', 'text', 'metric_refs', 'table_ids', 'chart_ids'}:
                raise StoreError('Invalid report section fields', 422, 'invalid_narrative')
            refs = []
            if not isinstance(section.get('metric_refs', []), list):
                raise StoreError('metric_refs must be a list', 422, 'invalid_narrative')
            for ref in section.get('metric_refs', []):
                if not isinstance(ref, dict) or set(ref) - {'analysis_id', 'metric_key', 'label'} or not {'analysis_id', 'metric_key'}.issubset(ref):
                    raise StoreError('Invalid metric reference', 422, 'invalid_narrative')
                refs.append(dict(analysis_id=ref['analysis_id'], key=ref['metric_key'],
                    label=self._prose(ref['label'], indexed) if ref.get('label') else self._clean(ref['metric_key']),
                    value=self._metric(ref['analysis_id'], ref['metric_key'], indexed)))
            for field, authorized in [('table_ids', tables), ('chart_ids', charts)]:
                if component_state is not None:
                    authorized = set(authorized) | {s['component_id'] for s in component_state['tables' if field == 'table_ids' else 'charts']}
                if not isinstance(section.get(field, []), list) or any(value not in authorized for value in section.get(field, [])):
                    raise StoreError('Section references an unauthorized table or chart', 422, 'artifact_not_authorized')
            sections.append(dict(heading=self._prose(section.get('heading', '分析结果'), indexed),
                text=self._prose(section.get('text', ''), indexed), metrics=refs,
                table_ids=section.get('table_ids', []), chart_ids=section.get('chart_ids', [])))
        extra_limits = narrative.get('limitations', [])
        if not isinstance(extra_limits, list):
            raise StoreError('limitations must be a list', 422, 'invalid_narrative')
        limitations = list(dict.fromkeys([self._clean(x) for a in analyses for x in a.limitations] +
            [self._prose(x, indexed) for x in extra_limits]))
        # Findings are tool-generated source facts, not model-authored prose.
        findings = list(dict.fromkeys(self._clean(x) for a in analyses if not any(ref.get('source_type') == 'material' for ref in a.kb_refs) for x in a.findings))
        summary = self._prose(narrative['summary'], indexed) if 'summary' in narrative else ''
        if not summary.strip() or summary.strip() == '本报告基于已验证分析结果整理，指标与来源可逐项追溯。':
            summary = self._default_summary(analyses, task.request)
        return dict(title=self._prose(title, indexed), version=version,
            generated_at=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds'), timezone='Asia/Shanghai',
            summary=summary,
            metrics=metrics, sections=sections, tables=list(tables.values()), charts=list(charts.values()),
            sources=list(sources.values()), quality=quality, findings=findings, limitations=limitations,
            limitations_fallback=True,
            materials=materials, html_table_limit=HTML_TABLE_LIMIT)

    def _render_patch_section(self, raw, previous, fields, indexed, state):
        refs = []
        if previous is not None and 'metric_refs' not in fields:
            refs = json.loads(json.dumps(previous['metrics']))
        else:
            for ref in raw.get('metric_refs', []):
                refs.append(dict(analysis_id=ref['analysis_id'], key=ref['metric_key'],
                    label=self._prose(ref['label'], indexed) if ref.get('label') else self._clean(ref['metric_key']),
                    value=self._metric(ref['analysis_id'], ref['metric_key'], indexed)))
        result = dict(metrics=refs)
        for field in ('heading', 'text'):
            result[field] = (previous[field] if previous is not None and field not in fields else
                             self._prose(raw.get(field, '分析结果' if field == 'heading' else ''), indexed))
        for field, key in [('table_ids', 'tables'), ('chart_ids', 'charts')]:
            authorized = {s[k] for s in state[key] for k in ('component_id', 'artifact_id')}
            values = previous[field] if previous is not None and field not in fields else raw.get(field, [])
            if any(value not in authorized for value in values):
                raise StoreError('Chapter artifact is not in the revised report', 422, 'artifact_not_authorized')
            result[field] = values
        return result

    @staticmethod
    def _revision_integrity(base, context, state):
        """Record actual component-level changes and unchanged content hashes."""
        from src.report_revision import digest
        before, after = {}, {}
        for field in ('title', 'summary', 'findings', 'limitations', 'quality'):
            before[field], after[field] = base['presentation'][field], context[field]
        before['limitations'] = dict(items=before['limitations'], fallback=base['presentation']['limitations_fallback'])
        after['limitations'] = dict(items=after['limitations'], fallback=context['limitations_fallback'])
        for option in context['layout']:
            before['layout:' + option] = base['presentation']['layout'][option]
            after['layout:' + option] = context['layout'][option]
        for section in base['presentation']['sections']:
            before.setdefault('section:' + section['heading'], []).append(section)
        for section in context['sections']:
            after.setdefault('section:' + section['heading'], []).append(section)
        for kind in ('charts', 'tables'):
            for slot in base['component_state'][kind]:
                before[kind + ':' + slot['component_id']] = slot
            for slot in state[kind]:
                after[kind + ':' + slot['component_id']] = slot
            before[kind + ':order'] = [s['component_id'] for s in base['component_state'][kind]]
            after[kind + ':order'] = [s['component_id'] for s in state[kind]]
        before['sections:order'] = [s['heading'] for s in base['presentation']['sections']]
        after['sections:order'] = [s['heading'] for s in context['sections']]
        keys = sorted(set(before) | set(after))
        changed = [key for key in keys if key not in before or key not in after or before[key] != after[key]]
        unchanged = {key: digest(before[key]) for key in keys if key not in changed}
        return dict(changed_targets=changed, unchanged_hashes=unchanged,
                    original_report_hash=base['hash'], source_lineage='append-only')

    def _validate_revision_scope(self, task, base, context, state, patch):
        """Fail closed when a concrete persisted edit plan is exceeded."""
        integrity = self._revision_integrity(base, context, state)
        spec = self.store.get_task(task.task_id).get('spec', {})
        plan = spec.get('constraints', {}).get('report_revision') or {}
        if plan:
            if plan.get('base_report_id') != base['report_id'] or plan.get('base_report_hash') != base['hash']:
                raise StoreError('Report revision does not match its pinned base', 409, 'revision_base_changed')
        targets = set(plan.get('targets', []))
        if not targets:
            return integrity  # Legacy direct tool use has no persisted plan.
        for target in targets:
            if target.startswith('section:') and sum(s['heading'] == target[8:] for s in base['presentation']['sections']) != 1:
                raise StoreError('Chapter target is ambiguous in the pinned report', 422, 'revision_target_not_found')
        permitted = set(targets)
        if patch is not None:
            requested = {field for field in ('summary', 'findings', 'limitations') if field in patch.model_fields_set}
            if patch.quality_edits:
                requested.add('quality')
            for edit in patch.components:
                plural = 'charts' if edit.kind == 'chart' else 'tables'
                if edit.operation == 'add':
                    requested.add(plural + ':add')
                else:
                    slot = next((s for s in base['component_state'][plural]
                        if edit.target_id in {s['component_id'], s['artifact_id']}), None)
                    requested.add(edit.kind + ':' + (slot['component_id'] if slot else str(edit.target_id)))
            requested.update('sections:add' if op.operation == 'add' else 'section:' + str(op.target_heading) for op in patch.sections)
            if requested - targets:
                raise StoreError('Patch operation exceeds the declared revision targets: ' + ', '.join(sorted(requested - targets)),
                                 422, 'revision_scope_exceeded')
        for key, singular in [('charts', 'chart'), ('tables', 'table')]:
            permitted.update(key + ':' + target.split(':', 1)[1] for target in targets if target.startswith(singular + ':'))
            if any(t.startswith(singular + ':') for t in targets) or key + ':add' in targets:
                permitted.add(key + ':order')
            if patch is not None and key + ':add' in targets:
                permitted.update(key + ':' + edit.artifact_id for edit in patch.components
                    if edit.kind == singular and edit.operation == 'add' and edit.artifact_id)
                permitted.add('layout:include_charts' if key == 'charts' else 'layout:include_tables')
        if any(target.startswith('section:') for target in targets) or 'sections:add' in targets:
            permitted.add('sections:order')
        for plural, singular in [('charts', 'chart'), ('tables', 'table')]:
            untouched = [s['component_id'] for s in base['component_state'][plural]
                         if singular + ':' + s['component_id'] not in targets]
            current_order = [s['component_id'] for s in state[plural] if s['component_id'] in untouched]
            if current_order != untouched:
                raise StoreError('Unspecified report components were removed or reordered', 422, 'revision_scope_exceeded')
        untouched_sections = [s['heading'] for s in base['presentation']['sections'] if 'section:' + s['heading'] not in targets]
        if [s['heading'] for s in context['sections'] if s['heading'] in untouched_sections] != untouched_sections:
            raise StoreError('Unspecified chapters were removed or reordered', 422, 'revision_scope_exceeded')
        if patch is not None:
            for op in patch.sections:
                if op.section and ('sections:add' in targets and op.operation == 'add' or
                        'section:' + str(op.target_heading) in targets and op.operation == 'update'):
                    permitted.add('section:' + op.section.heading)
        removed_headings = {s['heading'] for s in base['presentation']['sections']} - {s['heading'] for s in context['sections']}
        new_headings = {s['heading'] for s in context['sections']} - {s['heading'] for s in base['presentation']['sections']}
        # A legacy one-for-one heading translation/rename has an old stable
        # heading target. New sections require sections:add separately.
        if new_headings and len(new_headings) <= len(removed_headings) and all('section:' + h in targets for h in removed_headings):
            permitted.update('section:' + h for h in new_headings)
        # Removing a chart/table also clears that now-dead attachment reference
        # from a chapter; its text/metrics must remain exactly unchanged.
        for old in base['presentation']['sections']:
            current = next((s for s in context['sections'] if s['heading'] == old['heading']), None)
            if current and all(current[k] == old[k] for k in ('heading', 'text', 'metrics')):
                removed = set(old['chart_ids'] + old['table_ids']) - set(current['chart_ids'] + current['table_ids'])
                if removed and all('chart:' + aid in targets or 'table:' + aid in targets for aid in removed):
                    permitted.add('section:' + old['heading'])
        excess = set(integrity['changed_targets']) - permitted
        if excess:
            raise StoreError('Report revision exceeds the declared targets: ' + ', '.join(sorted(excess)),
                             422, 'revision_scope_exceeded')
        return integrity

    @staticmethod
    def _cell(sheet, row, column, value):
        # Workbook/cell public APIs, with explicit data_type='s' for every text
        # value (including =,+,-,@ prefixes and Excel error-like strings).
        # https://openpyxl.readthedocs.io/en/stable/tutorial.html
        # https://openpyxl.readthedocs.io/en/stable/api/openpyxl.cell.cell.html
        if isinstance(value, (dict, list)):
            value = exact(value)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            lexical = str(value)
            if (isinstance(value, float) and not math.isfinite(value)) or len(re.sub(r'[^0-9]', '', lexical.lstrip('-').replace('.', '').lstrip('0'))) > 15:
                value = lexical
        cell = sheet.cell(row=row, column=column, value=value)
        if isinstance(value, str):
            cell.data_type = 's'
            cell.number_format = '@'
        return cell

    def _sheet(self, workbook, name, rows):
        sheet = workbook.create_sheet(name)
        for row_number, row in enumerate(rows, 1):
            for col_number, value in enumerate(row, 1):
                self._cell(sheet, row_number, col_number, value)
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='24374B')
        for col in range(1, min(sheet.max_column, 30) + 1):
            sheet.column_dimensions[sheet.cell(1, col).column_letter].width = 28
        return sheet

    @staticmethod
    def _table_value(value):
        # Numeric CSV lexemes may be numbers; leading-zero IDs, NA and every
        # other text remain exact strings. High precision stays lexical.
        if re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?', value) and len(re.sub(r'[^0-9]', '', value).lstrip('0')) <= 15:
            number = float(value) if any(c in value for c in '.eE') else int(value)
            return number if not isinstance(number, float) or math.isfinite(number) else value
        return value

    def _xlsx(self, path, context):
        workbook = Workbook()
        workbook.remove(workbook.active)
        self._sheet(workbook, '说明与口径', [['项目', '内容'], ['标题', context['title']], ['报告版本', context['version']],
            ['生成时间', context['generated_at']], ['时区', context['timezone']],
            *([['结论', context['summary']]] if context['layout']['include_conclusion'] else []),
            ['数值精度', '指标精确值列保留原始表示；超过Excel有效位数的数字以文本保存。'],
            *([['文本处理', f'非法XML控制字符已移除：{self.cleaned_characters}；外来文本按字符串写入，未执行公式。']] if context['layout']['include_quality_and_limitations'] else []),
            ['明细范围', ('已按要求附上完整结果表；HTML长表仅显示明确标注的预览。'
                          if context['layout']['include_tables'] else '报告仅呈现结论与必要分析；完整结果保留为独立产物，可另行下载。')],
            *[['主要发现', x] for x in context['findings'] if context['layout']['include_findings']],
            *[['局限', x] for x in context['limitations'] if context['layout']['include_quality_and_limitations']]])
        metric_rows = [['analysis_id', 'metric_key', '指标', '精确值', '数值或原值', '口径']]
        metric_rows += [[m['analysis_id'], m['key'], m['label'], m['value'], m['raw'], m['definition']] for m in context['metrics']]
        if context['layout']['include_metrics_summary']:
            self._sheet(workbook, '指标汇总', metric_rows)
        for index, section in enumerate(context['sections'], 1):
            self._sheet(workbook, f'章节_{index}', [['项目', '内容'], ['标题', section['heading']], ['正文', section['text']],
                *[[m['label'], m['value']] for m in section['metrics']]])
        if context['materials']:
            self._sheet(workbook, '材料摘录', [['analysis_id', 'file_id', '定位', '原文摘录（非统计结论）', '范围说明'],
                *[[m['analysis_id'], m['file_id'], m['location'], m['text'],
                   '原文已截断；完整内容见源文件，不将摘录视为全文。' if m['truncated'] else '完整提取正文；未进行统计分析。'] for m in context['materials']]])
        for index, table in enumerate(context['tables'] if context['layout']['include_tables'] else [], 1):
            for offset in range(0, max(len(table['rows']), 1), 1048575):
                part = table['rows'][offset:offset + 1048575]
                self._sheet(workbook, f'结果表_{index}_{offset // 1048575 + 1}',
                    [table['columns'], *[[self._table_value(v) for v in row] for row in part]])
        if context['layout']['include_quality_and_limitations']:
            self._sheet(workbook, '数据质量', [['来源ID', '工作表或查询表', '说明'],
                *[[q.get('file_id', q.get('source_id', '')), q['table'], q['text']] for q in context['quality']]])
        if context['layout']['include_sources']:
            self._sheet(workbook, '数据来源', [['来源ID', '名称', '来源', 'SHA256'],
                *[[s.get('file_id', s.get('source_id', '')), s['name'], s['source'], s['hash']] for s in context['sources']]])
            database_sources = [s for s in context['sources'] if s.get('source_type') == 'postgresql']
            if database_sources:
                self._sheet(workbook, '数据库查询', [['source_id', 'query_id', '查询时间', '权限范围', '权限指纹',
                    'MVCC快照', '完整结果行数', 'SQL', '参数', '结果SHA256', 'SQL SHA256', '元数据SHA256'],
                    *[[s['source_id'], s['query_id'], s['queried_at'], s['scope'], s['scope_fingerprint'],
                        s['snapshot_id'], s['row_count'], s['sql'], s['params'], s['result_hash'],
                        s['sql_hash'], s['metadata_hash']] for s in database_sources]])
        if context['layout']['include_charts'] and context['charts']:
            chart_sheet = self._sheet(workbook, '图表', [['图表', '来源analysis_id']])
            for index, chart in enumerate(context['charts']):
                self._cell(chart_sheet, 2 + index * 28, 1, chart['name'])
                self._cell(chart_sheet, 2 + index * 28, 2, chart['analysis_id'])
                picture = SheetImage(io.BytesIO(chart['raw']))
                picture.width, picture.height = 720, 405
                chart_sheet.add_image(picture, f'A{3 + index * 28}')
        workbook.save(path)
        workbook.close()

    def _validate(self, staging, formats, context):
        if 'html' in formats:
            soup = BeautifulSoup((staging / 'report.html').read_text(encoding='utf-8'), 'html.parser')
            actual = {(node['data-analysis-id'], node['data-metric-key']): node.get_text() for node in soup.select('.metric-value')}
            for metric in context['metrics'] if context['layout']['include_metrics_summary'] else []:
                if actual.get((metric['analysis_id'], metric['key'])) != metric['value']:
                    raise StoreError('HTML metric validation failed', 422, 'report_validation_failed')
            if not context['layout']['include_metrics_summary'] and (actual or soup.find('h2', string='指标汇总')):
                raise StoreError('Removed HTML metrics summary is still present', 422, 'report_validation_failed')
            conclusion = soup.select_one('[data-section="conclusion"]')
            if context['layout']['include_conclusion'] and (not conclusion or conclusion.find_next_sibling('p').get_text() != context['summary']):
                raise StoreError('HTML conclusion validation failed', 422, 'report_validation_failed')
            if not context['layout']['include_conclusion'] and conclusion:
                raise StoreError('Removed HTML conclusion is still present', 422, 'report_validation_failed')
            if not context['layout']['include_tables'] and soup.find('h2', string='结果明细'):
                raise StoreError('Removed HTML result details are still present', 422, 'report_validation_failed')
            for option, heading in [('include_quality_and_limitations', '数据质量与局限'), ('include_sources', '来源与可追溯信息')]:
                if not context['layout'][option] and soup.find('h2', string=heading):
                    raise StoreError('Removed HTML section is still present: ' + heading, 422, 'report_validation_failed')
            scripts = soup.find_all('script')
            expected_charts = context['charts'] if context['layout']['include_charts'] else []
            if (len(scripts) != (1 if expected_charts else 0)
                    or any(script.get_text() != interaction_script() or set(script.attrs) != {'id'} or script.get('id') != 'ds-chart-interactions' for script in scripts)
                    or soup.find(['img', 'iframe', 'object', 'embed', 'link'])
                    or len(soup.select('svg[data-ds-chart]')) != len(expected_charts)
                    or any(key.lower().startswith('on') for node in soup.find_all(True) for key in node.attrs)):
                raise StoreError('HTML offline content validation failed', 422, 'report_validation_failed')
            if soup.find('meta', attrs={'http-equiv': 'Content-Security-Policy'}).get('content') != chart_csp():
                raise StoreError('HTML chart CSP validation failed', 422, 'report_validation_failed')
        if 'xlsx' in formats:
            book = load_workbook(staging / 'report.xlsx', read_only=True, data_only=False)
            try:
                if context['layout']['include_conclusion'] and (dict(book['说明与口径'].values).get('结论') or '') != context['summary']:
                    raise StoreError('XLSX conclusion validation failed', 422, 'report_validation_failed')
                if not context['layout']['include_conclusion'] and '结论' in dict(book['说明与口径'].values):
                    raise StoreError('Removed XLSX conclusion is still present', 422, 'report_validation_failed')
                if not context['layout']['include_tables'] and any(name.startswith('结果表_') for name in book.sheetnames):
                    raise StoreError('Removed XLSX result details are still present', 422, 'report_validation_failed')
                if context['layout']['include_metrics_summary']:
                    rows = list(book['指标汇总'].values)[1:]
                    if [row[3] for row in rows] != [m['value'] for m in context['metrics']]:
                        raise StoreError('XLSX metric validation failed', 422, 'report_validation_failed')
                elif '指标汇总' in book.sheetnames:
                    raise StoreError('Removed XLSX metrics summary is still present', 422, 'report_validation_failed')
                for option, sheet in [('include_quality_and_limitations', '数据质量'), ('include_sources', '数据来源')]:
                    if not context['layout'][option] and sheet in book.sheetnames:
                        raise StoreError('Removed XLSX section is still present: ' + sheet, 422, 'report_validation_failed')
                if not context['layout']['include_quality_and_limitations'] and any(row[0] == '局限' for row in book['说明与口径'].values):
                    raise StoreError('Removed XLSX limitations are still present', 422, 'report_validation_failed')
                for sheet in book:
                    if any(cell.data_type == 'f' for row in sheet for cell in row):
                        raise StoreError('XLSX contains an untrusted formula', 422, 'report_validation_failed')
                for index, table in enumerate(context['tables'] if context['layout']['include_tables'] else [], 1):
                    count = sum(book[name].max_row - 1 for name in book.sheetnames if name.startswith(f'结果表_{index}_'))
                    if count != table['total_rows']:
                        raise StoreError('XLSX result row validation failed', 422, 'report_validation_failed')
            finally:
                book.close()
        for format in formats:
            if not (staging / f'report.{format}').stat().st_size:
                raise StoreError('Empty report output', 422, 'report_validation_failed')

    @traced_method('report_generate')
    @validated_method
    def generate(self, task: TaskSpec, analysis_ids: list[str], title: str | None = None, narrative: dict | None = None,
                 formats: list[str] | None = None, *, layout: dict | None = None, base_report_id: str | None = None,
                 text_replacements: list[dict] | None = None,
                 chart_changes: list[dict] | None = None, section_edits: list[dict] | None = None,
                 revision_patch: dict | None = None) -> ReportResult:
        if layout is not None and not isinstance(layout, dict):
            raise StoreError('Invalid report layout options', 422, 'invalid_layout')
        base = self.read_report(task, base_report_id) if base_report_id else None
        if text_replacements is not None and (not base or not isinstance(text_replacements, list) or len(text_replacements) > 300):
            raise StoreError('翻译替换需要已有报告，最多允许300条。', 422, 'invalid_text_replacement')
        from src.report_revision import (initial_state, validate_state, apply_components,
                                         explicit_analysis_ids, clone, digest)
        try:
            patch = ReportRevisionPatch.model_validate(revision_patch) if revision_patch is not None else None
        except (ValidationError, TypeError) as exc:
            raise StoreError('Invalid general report revision patch', 422, 'invalid_report_edit') from exc
        if patch is not None and not base:
            raise StoreError('A revision patch requires a delivered base report', 422, 'missing_report')
        if base:
            old_ids, all_ids = set(base['analysis_ids']), set(analysis_ids)
            added_ids = all_ids - old_ids
            if not old_ids.issubset(all_ids) or (added_ids and
                    (patch is None or not added_ids.issubset(explicit_analysis_ids(patch,
                        {**(narrative or {}), **({'title': title} if title is not None else {})})))):
                raise StoreError('Revision must preserve all source analyses and explicitly bind supplemental analysis',
                                 422, 'report_source_mismatch')
        try:
            if (chart_changes is not None or section_edits is not None) and not base:
                raise ValueError('局部修改需要已有报告。')
            if len(chart_changes or []) > 30 or len(section_edits or []) > 100:
                raise ValueError('局部修改项目过多。')
            chart_edits = [ReportChartEdit.model_validate(item) for item in chart_changes or []]
            chapter_edits = [ReportSectionEdit.model_validate(item) for item in section_edits or []]
            if len({item.chart_id for item in chart_edits}) != len(chart_edits):
                raise ValueError('同一图表只能指定一次修改。')
            if len({item.section_heading for item in chapter_edits}) != len(chapter_edits):
                raise ValueError('同一章节只能指定一次修改。')
        except (ValidationError, ValueError, TypeError) as exc:
            raise StoreError('局部修改参数无效：' + str(exc), 422, 'invalid_report_edit') from exc
        try:
            # Every new deliverable, including a revision of a legacy report,
            # omits the metrics summary unless explicitly requested. Other
            # visibility switches continue to inherit the base version.
            presentation_layout = ReportLayout.model_validate({**(base['presentation']['layout'] if base else {}),
                **({'include_metrics_summary': False} if not base or (patch is None and not task.constraints.get('report_revision')) else {}),
                **(layout or {})}).model_dump()
        except (ValidationError, TypeError) as exc:
            raise StoreError('Invalid report layout options', 422, 'invalid_layout') from exc
        chosen = list(dict.fromkeys(formats if formats is not None else (task.expected_outputs or (base['formats'] if base else ['html', 'xlsx']))))
        if not chosen or any(x not in {'html', 'xlsx'} for x in chosen):
            raise StoreError('Choose html and/or xlsx report formats', 422, 'invalid_format')
        if not analysis_ids:
            return ReportResult(status='need_analysis', validation_summary={'reason_code': 'missing_analysis'},
                                supplement_request='请先提供带来源的已验证分析结果或材料摘录。')
        self.cleaned_characters = 0
        analyses = self._analyses(task, analysis_ids)
        if base:
            for analysis in analyses:
                if analysis.analysis_id in added_ids:
                    origin = self.store.get_task(analysis.task_id)
                    if origin['status'] != 'succeeded' and analysis.task_id != task.task_id:
                        raise StoreError('Supplemental analysis was not successfully delivered', 422, 'analysis_not_authorized')
        component_state = clone(base['component_state']) if base else None
        component_changes = []
        if component_state is not None:
            # A whole-module visibility request is explicit. A newly added
            # table only turns on its own slot, never all old hidden CSVs.
            for kind, option in [('charts', 'include_charts'), ('tables', 'include_tables')]:
                if option in (layout or {}):
                    for slot in component_state[kind]:
                        slot['visible'] = presentation_layout[option]
            if patch is not None:
                component_state, component_changes = apply_components(self.store, task, analyses, component_state, patch.components)
            validate_state(self.store, task, analyses, component_state, base.get('chart_overrides', []))
            for kind, option in [('charts', 'include_charts'), ('tables', 'include_tables')]:
                if option not in (layout or {}) and any(s['visible'] for s in component_state[kind]):
                    presentation_layout[option] = True
        self.require_svg = 'html' in chosen and presentation_layout['include_charts']
        self.require_png = 'xlsx' in chosen and presentation_layout['include_charts']
        self.include_result_tables = presentation_layout['include_tables']
        version = datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%dT%H%M%S') + '_' + new_id()
        try:
            context = self._context(task, analyses, title or '数据分析报告', narrative, version, component_state)
        except NeedAnalysis as exc:
            return ReportResult(status='need_analysis', analysis_ids=[a.analysis_id for a in analyses],
                                validation_summary={'reason_code': exc.reason_code}, supplement_request=str(exc))
        context['layout'] = presentation_layout
        context['chart_script'] = interaction_script()
        context['chart_csp'] = chart_csp()
        context['report_css'] = (Path(__file__).resolve().parents[1] / 'templates' / 'report.css').read_text(encoding='utf-8')
        source_quality = [dict(item) for item in context['quality']]
        if base:
            if title is None:
                context['title'] = base['presentation']['title']
            for key in ('summary', 'sections', 'limitations'):
                if key not in (narrative or {}):
                    context[key] = json.loads(json.dumps(base['presentation'].get(key, [])))
            if narrative is not None and 'summary' in narrative:
                context['summary'] = self._prose(narrative['summary'], {a.analysis_id: a for a in analyses})
            if narrative is not None and 'limitations' in narrative:
                # An explicit revised list REPLACES the displayed list. Raw
                # analysis limitations remain immutable in manifest.analyses.
                context['limitations'] = list(dict.fromkeys(self._prose(x, {a.analysis_id: a for a in analyses}) for x in narrative['limitations']))
            context['findings'] = list(base['presentation']['findings'])
            context['quality'] = [dict(item) for item in base['presentation']['quality']]
            context['limitations_fallback'] = base['presentation']['limitations_fallback']
            if narrative is not None and 'limitations' in narrative:
                context['limitations_fallback'] = False
            # Supplemental analysis does not inject its material excerpts.
            context['materials'] = [m for m in context['materials'] if m['analysis_id'] in base['analysis_ids']]
            for key, state_key in [('chart_ids', 'charts'), ('table_ids', 'tables')]:
                active = {s[field] for s in component_state[state_key] for field in ('component_id', 'artifact_id')}
                for section in context['sections']:
                    section[key] = [aid for aid in section[key] if aid in active]
        from src.report_edits import plan_edit, discard_artifacts
        chart_overrides = {item['chart_id']: dict(item) for item in (base.get('chart_overrides', []) if base else [])}
        if component_state is None:
            component_state = initial_state(self.store, analyses, presentation_layout)
        # Reuse inherited overrides even for an unrelated subsequent edit;
        # never restore the first report's TopN or translated prose.
        for index, chart in enumerate(context['charts']):
            override = chart_overrides.get(chart['artifact_id'])
            if override:
                context['charts'][index] = self._chart(override['artifact_id'], override['analysis_id'])
        plans = []
        for edit in chart_edits:
            active = next((s for s in component_state['charts'] if s['component_id'] == edit.chart_id), None)
            if active is None or not active['visible']:
                raise StoreError('目标图表不是本版本中可见的图表。', 422, 'artifact_not_authorized')
            if active['artifact_id'] not in {edit.chart_id, chart_overrides.get(edit.chart_id, {}).get('artifact_id')}:
                return ReportResult(status='need_analysis', analysis_ids=analysis_ids,
                    validation_summary={'reason_code': 'chart_edit_source_required', 'target_chart_id': edit.chart_id},
                    supplement_request='此图已被新的分析图替换。请基于当前图及其已登记数据，只重绘该目标图并通过revision_patch替换。')
            owner = next((a for a in analyses if edit.chart_id in a.chart_ids), None)
            if owner is None:
                raise StoreError('目标图表不属于原报告。', 422, 'artifact_not_authorized')
            try:
                spec = plan_edit(self.store, edit.chart_id, [owner], edit, chart_overrides.get(edit.chart_id))
            except StoreError as exc:
                if exc.code == 'chart_edit_source_required':
                    return ReportResult(status='need_analysis', analysis_ids=analysis_ids,
                        validation_summary={'reason_code': exc.code, 'target_chart_id': edit.chart_id}, supplement_request=str(exc))
                raise
            if spec:
                plans.append((edit.chart_id, owner.analysis_id, spec))
        changed_chapters = []
        indexed = {a.analysis_id: a for a in analyses}
        if patch is not None:
            try:
                if len({edit.index for edit in patch.quality_edits}) != len(patch.quality_edits):
                    raise StoreError('A quality note may only be edited once', 422, 'invalid_report_edit')
                for edit in sorted(patch.quality_edits, key=lambda item: item.index, reverse=True):
                    if edit.index >= len(context['quality']):
                        raise StoreError('Quality note index does not exist in the base report', 422, 'quality_not_found')
                    if edit.remove:
                        if edit.text is not None:
                            raise StoreError('Removing a quality note cannot also replace its text', 422, 'invalid_report_edit')
                        context['quality'].pop(edit.index)
                    elif edit.text is not None:
                        context['quality'][edit.index]['text'] = self._prose(edit.text, indexed)
                    else:
                        raise StoreError('Quality edit requires remove or text', 422, 'invalid_report_edit')
                for field in ('summary', 'findings', 'limitations'):
                    if field in patch.model_fields_set:
                        value = getattr(patch, field)
                        if value is None:
                            raise StoreError('Explicit null is not a report text edit; omit it to inherit', 422, 'invalid_report_edit')
                        context[field] = (self._prose(value, indexed) if field == 'summary' else
                                          [self._prose(item, indexed) for item in value])
                        if field == 'limitations':
                            context['limitations_fallback'] = False
                for operation in patch.sections:
                    sections = context['sections']
                    if operation.operation == 'add':
                        if operation.target_heading is not None or operation.section is None:
                            raise StoreError('Adding a chapter requires only its new section', 422, 'invalid_report_edit')
                        previous = None
                    else:
                        matches = [item for item in sections if item['heading'] == operation.target_heading]
                        if len(matches) != 1:
                            raise StoreError('Chapter target is not unique; read the base report first', 422, 'section_not_found')
                        previous = matches[0]
                    if operation.operation in {'add', 'update'}:
                        if operation.section is None:
                            raise StoreError('Chapter content is required', 422, 'invalid_report_edit')
                        raw = operation.section.model_dump(exclude_unset=True)
                        if previous:
                            raw = {**dict(heading=previous['heading'], text=previous['text'],
                                metric_refs=[dict(analysis_id=m['analysis_id'], metric_key=m['key'], label=m['label']) for m in previous['metrics']],
                                table_ids=previous['table_ids'], chart_ids=previous['chart_ids']), **raw}
                        # Preserve unchanged rendered text, including its bound
                        # numbers; only new fields pass through the prose guard.
                        rendered = self._render_patch_section(raw, previous, operation.section.model_fields_set, indexed, component_state)
                        if previous:
                            index = sections.index(previous); sections[index] = rendered
                        else:
                            if any(s['heading'] == rendered['heading'] for s in sections):
                                raise StoreError('New chapter heading is ambiguous', 422, 'invalid_report_edit')
                            sections.append(rendered)
                        if operation.position is not None:
                            sections.remove(rendered)
                            if operation.position > len(sections):
                                raise StoreError('Chapter position exceeds report bounds', 422, 'invalid_report_edit')
                            sections.insert(operation.position, rendered)
                    else:
                        if operation.section is not None:
                            raise StoreError('Remove/move cannot replace chapter content', 422, 'invalid_report_edit')
                        if operation.operation == 'remove' and operation.position is not None:
                            raise StoreError('Removing a chapter must not specify a new position', 422, 'invalid_report_edit')
                        sections.remove(previous)
                        if operation.operation == 'move':
                            if operation.position is None or operation.position > len(sections):
                                raise StoreError('Move requires a valid chapter position', 422, 'invalid_report_edit')
                            sections.insert(operation.position, previous)
                    changed_chapters.append(dict(section_heading=operation.target_heading or rendered['heading'], operation=operation.operation))
                if len(context['sections']) > 100:
                    raise StoreError('Revised report exceeds the chapter limit', 422, 'invalid_report_edit')
            except NeedAnalysis as exc:
                return ReportResult(status='need_analysis', analysis_ids=analysis_ids,
                    validation_summary={'reason_code': exc.reason_code}, supplement_request=str(exc))
        for edit in chapter_edits:
            matches = [section for section in context['sections'] if section['heading'] == edit.section_heading]
            if len(matches) != 1:
                raise StoreError('章节标题未唯一匹配，请先read_report并使用完整标题。', 422, 'section_not_found')
            section = matches[0]
            if edit.remove:
                if edit.heading is not None or edit.text is not None:
                    raise StoreError('删除章节不能同时指定新标题或正文。', 422, 'invalid_report_edit')
                context['sections'].remove(section)
                changed_chapters.append(dict(section_heading=edit.section_heading, operation='remove'))
            else:
                previous = dict(section)
                try:
                    if edit.heading is not None:
                        section['heading'] = self._prose(edit.heading, indexed)
                    if edit.text is not None:
                        section['text'] = self._prose(edit.text, indexed)
                except NeedAnalysis as exc:
                    return ReportResult(status='need_analysis', analysis_ids=analysis_ids,
                        validation_summary={'reason_code': exc.reason_code}, supplement_request=str(exc))
                if section != previous:
                    changed_chapters.append(dict(section_heading=edit.section_heading, operation='update'))
        text_changes = self._replace_text(context, text_replacements or [])
        layout_changes = {key: value for key, value in presentation_layout.items()
                          if base and value != base['presentation']['layout'][key]}
        visual_upgrade = bool(base and base['needs_visual_upgrade'] and 'html' in chosen)
        if base and component_state == base['component_state'] and not plans and not visual_upgrade and chosen == base['formats'] and all(context[key] == base['presentation'].get(key, [])
                for key in ('title', 'summary', 'sections', 'layout', 'limitations', 'findings', 'quality', 'limitations_fallback')):
            raise StoreError('修改未改变报告内容，未生成重复文件；请使用对应的章节显示开关或提供具体修改。', 422, 'report_unchanged')
        # Snapshot all references; revalidate after rendering and before publish.
        staging = self.store.safe_path(f'script_output/{task.session_id}/{task.task_id}/.report_staging_{version}')
        final = self.store.safe_path(f'report/{task.session_id}/{task.task_id}/{version}')
        registered = []
        derived_ids, derived_dirs, chart_edit_proofs = [], [], []
        staging.mkdir(parents=True)
        try:
            from src.report_edits import render_chart, validate_override
            for cid, aid, spec in plans:
                directory = self.store.safe_path(f'script_output/{task.session_id}/{task.task_id}/report_edit_{new_id()}')
                derived_dirs.append(directory)
                previous = chart_overrides.get(cid)
                override = render_chart(self.store, task, cid, aid, spec, directory)
                derived_ids.extend([override['artifact_id'], override['preview_artifact_id']])
                validate_override(self.store, task, analyses, override)
                current_id = previous['artifact_id'] if previous else cid
                index = next(i for i, chart in enumerate(context['charts']) if chart['artifact_id'] == current_id)
                context['charts'][index] = self._chart(override['artifact_id'], aid)
                chart_overrides[cid] = override
                slot = next(s for s in component_state['charts'] if s['component_id'] == cid)
                slot.update(artifact_id=override['artifact_id'], hash=override['artifact_hash'], analysis_id=aid)
                chart_edit_proofs.append(dict(chart_id=cid, artifact_id=override['artifact_id'], top_n=spec['top_n'],
                    title=spec['title'], source_table_id=spec['table_id'], business_queries_rerun=False))
            validate_state(self.store, task, analyses, component_state, list(chart_overrides.values()))
            integrity = self._validate_revision_scope(task, base, context, component_state, patch) if base else None
            if 'html' in chosen:
                with span(self.store,task.task_id,'report_html',format='html'):
                    rendered = self.environment.get_template('report.html').render(**context, cleaned_characters=self.cleaned_characters)
                    (staging / 'report.html').write_text(rendered, encoding='utf-8')
            if 'xlsx' in chosen:
                with span(self.store,task.task_id,'report_xlsx',format='xlsx'):
                    self._xlsx(staging / 'report.xlsx', context)
            with span(self.store,task.task_id,'report_validate'):
                self._validate(staging, chosen, context)
            # Rendering/plotting can be long. Final authorization must not use
            # the catalog memo from the tool's earlier source-read boundary.
            with span(self.store,task.task_id,'report_revalidate'), validation_scope(self.store,task,fresh=True):
                rechecked = self._analyses(task, analysis_ids)
                if [a.model_dump() for a in rechecked] != [a.model_dump() for a in analyses]:
                    raise StoreError('Analysis changed during report rendering', 409, 'analysis_changed')
                if base:
                    self.store.resolve_artifact(base['report_id'])
            artifact_refs = [self.store.get_artifact(x).model_dump() for a in analyses for x in a.table_ids + a.chart_ids + a.script_ids]
            for analysis in analyses:
                for snapshot in self._database_snapshots(task, analysis):
                    for key in ('result_artifact_id', 'sql_artifact_id', 'metadata_artifact_id'):
                        ref = snapshot[key]
                        if not any(a['artifact_id'] == ref for a in artifact_refs):
                            self.store.resolve_artifact(ref)
                            artifact_refs.append(self.store.get_artifact(ref).model_dump())
            for chart in context['charts']:
                for key in ('svg_artifact_id', 'preview_artifact_id'):
                    ref = chart[key]
                    if ref and not any(a['artifact_id'] == ref for a in artifact_refs):
                        self.store.resolve_artifact(ref)
                        artifact_refs.append(self.store.get_artifact(ref).model_dump())
            manifest = dict(schema_version=VERSION, report_version=version, task_id=task.task_id, session_id=task.session_id,
                generated_at=context['generated_at'], timezone='Asia/Shanghai', formats=chosen,
                models={role: getattr(self.store.settings, field) for role, field in
                    [('main', 'main_model'), ('analyst', 'analyst_model'), ('report_designer', 'report_model')]},
                sources=context['sources'], analyses=[a.model_dump() for a in analyses], source_artifacts=artifact_refs,
                metric_bindings=[{k: v for k, v in m.items() if k != 'raw'} for m in context['metrics']],
                tables=[{k: v for k, v in t.items() if k not in {'rows', 'columns'}} for t in context['tables']],
                missing_value_handling='See each analysis.metric_definitions; no report recalculation or imputation.',
                limitations=context['limitations'],
                presentation={key: context[key] for key in ('title', 'summary', 'sections', 'layout', 'limitations', 'findings', 'quality', 'limitations_fallback')},
                source_quality=source_quality, text_replacements=text_changes,
                chart_overrides=list(chart_overrides.values()), chart_changes=chart_edit_proofs, section_edits=changed_chapters,
                component_state=component_state,
                revision_patch=patch.model_dump(exclude_unset=True) if patch is not None else None,
                component_changes=component_changes,
                revision_integrity=integrity,
                chart_rendering=CHART_RENDERING, style_revision=REPORT_STYLE_REVISION,
                revision_of={key: base[key] for key in ('report_id', 'task_id', 'version', 'hash')} if base else None,
                text_sanitization=dict(removed_xml_characters=self.cleaned_characters, formulas='all external text stored explicitly as strings'),
                validation=dict(status='passed', metrics_checked=len(context['metrics']), full_result_rows_checked=presentation_layout['include_tables'],
                    conclusion_included=presentation_layout['include_conclusion'], result_details_included=presentation_layout['include_tables'],
                    metrics_summary_included=presentation_layout['include_metrics_summary'],
                    visible_metrics_checked=len(context['metrics']) if presentation_layout['include_metrics_summary'] else 0,
                    html_offline=True, artifacts_hash_checked=True,
                    html_chart_format='svg' if 'html' in chosen else None,
                    visual_upgrade=visual_upgrade,
                    interactive_charts=bool('html' in chosen and presentation_layout['include_charts'] and context['charts'])),
                files=[dict(name=f'report.{fmt}', sha256=file_hash(staging / f'report.{fmt}')) for fmt in chosen])
            (staging / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, allow_nan=False, indent=2), encoding='utf-8')
            final.parent.mkdir(parents=True, exist_ok=True)
            staging.rename(final)
            with self.store.lock:
                try:
                    for fmt in chosen:
                        registered.append(self.store.register_artifact(final / f'report.{fmt}', task.task_id, task.session_id,
                            'report', [a.analysis_id for a in analyses]))
                    manifest_artifact = self.store.register_artifact(final / 'manifest.json', task.task_id, task.session_id,
                            'report_manifest', [a.analysis_id for a in analyses])
                    registered.append(manifest_artifact)
                except BaseException:
                    index = self.store._index('artifacts')
                    for artifact in registered:
                        index.pop(artifact.artifact_id, None)
                    self.store._save_index('artifacts', index)
                    raise
            return ReportResult(report_ids=[a.artifact_id for a in registered if a.kind == 'report'],
                analysis_ids=[a.analysis_id for a in analyses], limitations=context['limitations'],
                validation_summary={**manifest['validation'], 'manifest_id': manifest_artifact.artifact_id,
                    'version': version, 'formats': chosen, 'removed_xml_characters': self.cleaned_characters,
                    'layout': presentation_layout, 'layout_changes': layout_changes,
                    'revision_of': manifest['revision_of'],
                    'chart_changes': chart_edit_proofs, 'section_edits': changed_chapters,
                    'component_changes': component_changes,
                    'supplemental_analysis_ids': sorted(added_ids) if base else [],
                    'revision_integrity': manifest['revision_integrity'],
                    'data_snapshots_reused': bool(plans),
                    'text_replacements_applied': len(text_changes),
                    'translated_fields': sorted({item['field'] for item in text_changes}),
                    'title_changed': bool(base and context['title'] != base['presentation']['title']),
                    'narrative_changed': bool(base and (context['summary'] != base['presentation']['summary'] or
                                                       context['sections'] != base['presentation']['sections'] or
                                                       any(context[key] != base['presentation'][key] for key in ('limitations', 'findings', 'quality', 'limitations_fallback'))))})
        except BaseException:
            discard_artifacts(self.store, derived_ids)
            for directory in derived_dirs:
                if directory.exists():
                    shutil.rmtree(directory)
            # Both paths are unique task-scoped directories returned by safe_path.
            for directory in (staging, final):
                if directory.exists():
                    shutil.rmtree(directory)
            raise
