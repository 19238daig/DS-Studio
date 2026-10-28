"""Local, deterministic data tools. Model expressions are never evaluated here."""
from __future__ import annotations

import csv
import io
import json
import zipfile
from collections import OrderedDict
from pathlib import Path
from threading import RLock
from typing import Any

import pandas as pd
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from src.schemas import AnalysisResult, DatasetProfile, TaskSpec, new_id
from src.storage import Store, StoreError
from src.performance import stage, traced_method


def clean_json(value):
    """Normalize pandas scalars, dates and missing values for strict JSON."""
    return json.loads(pd.DataFrame([{'value': value}]).to_json(orient='records', date_format='iso'))[0]['value']


def records(frame: pd.DataFrame) -> list[dict]:
    return json.loads(frame.to_json(orient='records', date_format='iso'))


def infer_numbers(frame: pd.DataFrame) -> pd.DataFrame:
    for column in frame.columns:
        series = frame[column]
        present = series.dropna().astype(str)
        # Leading zeros, text IDs and mixed columns retain their exact spelling.
        if present.empty or present.str.match(r'^[+-]?0\d+').any():
            continue
        try:
            frame[column] = pd.to_numeric(series, errors='raise')
        except (ValueError, TypeError):
            pass
    return frame


def validate_columns(frame: pd.DataFrame) -> pd.DataFrame:
    names = [' / '.join(str(x) for x in c if not str(x).startswith('Unnamed:')) if isinstance(c, tuple) else str(c) for c in frame.columns]
    if len(set(names)) != len(names) or any(not c.strip() for c in names):
        raise StoreError('Table has duplicate or empty column names; correct the source headers', 422, 'invalid_columns')
    frame.columns = names
    return frame


class DataService:
    def __init__(self, store: Store):
        self.store = store
        # Only this service/task instance owns these derived frames. Every
        # reuse still hashes the current original file; mtime/size is not an
        # authorization or integrity substitute.
        self._parsed = OrderedDict()
        self._parsed_bytes = 0
        # Tool batches may read the same task-owned service concurrently. Only
        # the small LRU operations are serialized; hashing/parsing stays outside.
        self._parsed_lock = RLock()
        self._result_shapes = {}

    def _load(self, file_id: str, session_id: str) -> tuple[dict[str, pd.DataFrame], list[str], str]:
        path = self.store.resolve_file(file_id, session_id)
        source = self.store.get_file(file_id, session_id)
        key = (session_id, file_id, source['hash'])
        with self._parsed_lock:
            if key in self._parsed:
                value, _ = self._parsed[key]
                self._parsed.move_to_end(key)
                return value
        with stage(self.store,'dataset_parse',file_count=1):
            value = self._load_uncached(file_id, session_id, path)
        # A concurrently changed original cannot seed the parsed-data cache.
        self.store.resolve_file(file_id, session_id)
        size = sum(int(frame.memory_usage(index=True, deep=True).sum()) for frame in value[0].values()) + len(value[2].encode('utf-8'))
        if size <= 64 * 1024 * 1024:
            with self._parsed_lock:
                # Another reader may have finished this exact parse meanwhile.
                # Keep its immutable derived result without counting it twice.
                if key in self._parsed:
                    cached, _ = self._parsed[key]
                    self._parsed.move_to_end(key)
                    return cached
                self._parsed[key] = (value, size)
                self._parsed_bytes += size
                while len(self._parsed) > 3 or self._parsed_bytes > 64 * 1024 * 1024:
                    _, (_, removed) = self._parsed.popitem(last=False)
                    self._parsed_bytes -= removed
        return value

    def _load_uncached(self, file_id: str, session_id: str, path: Path) -> tuple[dict[str, pd.DataFrame], list[str], str]:
        if path.stat().st_size > self.store.settings.max_upload_bytes:
            raise StoreError('Dataset exceeds configured file-size limit; no partial analysis performed', 422, 'input_too_large')
        suffix = path.suffix.lower()
        warnings: list[str] = []
        text = ''
        if suffix == '.csv':
            raw = path.read_bytes()
            for encoding in ('utf-8-sig', 'gb18030'):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise StoreError('CSV encoding must be UTF-8 or GB18030', 422, 'encoding_error')
            try:
                dialect = csv.Sniffer().sniff(text[:65536], delimiters=',;\t|')
                delimiter = dialect.delimiter
            except csv.Error:
                delimiter = ';' if text.splitlines() and ';' in text.splitlines()[0] else ','
            try:
                rows = list(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
            except csv.Error as exc:
                raise StoreError(f'Malformed CSV: {exc}', 422, 'malformed_csv') from exc
            rows = [row for row in rows if row]  # Match ordinary CSV blank-line semantics.
            if not rows:
                raise StoreError('CSV is empty', 422, 'empty_data')
            width = len(rows[0])
            bad = [i + 2 for i, row in enumerate(rows[1:]) if len(row) != width]
            if bad:
                raise StoreError(f'CSV row width differs from header at record(s): {bad[:10]}', 422, 'malformed_csv')
            frame = pd.DataFrame(rows[1:], columns=rows[0]).replace('', None)
            tables = {'data': infer_numbers(validate_columns(frame))}
            warnings.append(f'CSV encoding={encoding}; delimiter={repr(delimiter)}; all records parsed')
        elif suffix in {'.xlsx', '.xls'}:
            if suffix == '.xlsx':
                with zipfile.ZipFile(path) as archive:
                    entries = archive.infolist()
                    if len(entries) > 10000 or sum(item.file_size for item in entries) > self.store.settings.max_task_input_bytes:
                        raise StoreError('Expanded XLSX exceeds file count or uncompressed size limit', 422, 'input_too_large')
            tables = pd.read_excel(path, sheet_name=None, dtype=object, keep_default_na=False)
            tables = {name: infer_numbers(validate_columns(frame.replace('', None))) for name, frame in tables.items()}
            if suffix == '.xlsx':
                formula_book = load_workbook(path, read_only=True, data_only=False)
                cache_book = load_workbook(path, read_only=True, data_only=True)
                try:
                    formulas, absent = 0, 0
                    for sheet in formula_book:
                        cached = cache_book[sheet.title]
                        for row in sheet:
                            for cell in row:
                                if cell.data_type == 'f':
                                    formulas += 1
                                    absent += cached[cell.coordinate].value is None
                    if formulas:
                        warnings.append(f'{formulas} formula cell(s): cached values may be stale; formulas are not recalculated')
                    if absent:
                        warnings.append(f'{absent} formula cell(s) lack cached values and are read as missing')
                finally:
                    formula_book.close()
                    cache_book.close()
            else:
                warnings.append('Legacy XLS formulas are read from cached values; recalculation and cache freshness cannot be verified')
        elif suffix in {'.md', '.txt'}:
            # Treat documents as literal source text, never as Markdown/HTML to
            # render, executable instructions or synthetic tabular data.
            raw = path.read_bytes()
            for encoding in ('utf-8-sig', 'gb18030'):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise StoreError('Document encoding must be UTF-8 or GB18030', 422, 'encoding_error')
            if '\x00' in text:
                raise StoreError('Document contains NUL bytes; provide a plain-text document', 422, 'invalid_text')
            tables = {}
            warnings.append(f'Plain-text document encoding={encoding}; no tables inferred; markup and code were not executed')
        elif suffix in {'.html', '.htm'}:
            # Parse local bytes only: no browser, JavaScript, external URLs or resources.
            raw = path.read_bytes()
            try:
                html = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                html = raw.decode('gb18030')
            soup = BeautifulSoup(html, 'html.parser')
            for node in soup(['script', 'style', 'iframe', 'object', 'link', 'img']):
                node.decompose()
            text = soup.get_text(' ', strip=True)
            tables = {}
            for index, element in enumerate(soup.find_all('table')):
                if element.find_parent('table'):
                    continue
                try:
                    width = len(pd.read_html(io.StringIO(str(element)), keep_default_na=False)[0].columns)
                    frame = pd.read_html(io.StringIO(str(element)), converters={i: str for i in range(width)}, keep_default_na=False)[0]
                except ValueError:
                    continue
                tables[f'table_{index + 1}'] = infer_numbers(validate_columns(frame.replace('', None)))
            warnings.append('HTML parsed statically; scripts and external resources were not executed or fetched')
        else:
            raise StoreError('Unsupported dataset type', 422, 'unsupported_type')
        if any(len(frame) > 1000000 or frame.size > 5000000 for frame in tables.values()):
            raise StoreError('Table exceeds one million rows or five million cells; no partial analysis performed', 422, 'input_too_large')
        return tables, warnings, text

    def read_table(self, file_id: str, session_id: str, table: str | None = None) -> pd.DataFrame:
        tables, _, _ = self._load(file_id, session_id)
        if not tables:
            raise StoreError('No tables found; use read_text for document text', 422, 'no_tables')
        if table is None:
            if len(tables) != 1:
                raise StoreError(f'Select a sheet/table explicitly: {list(tables)}', 422, 'table_required')
            table = next(iter(tables))
        if table not in tables:
            raise StoreError(f'Unknown sheet/table {table!r}; available: {list(tables)}', 422, 'unknown_table')
        return tables[table].copy()

    def profiles(self, file_id: str, session_id: str) -> list[DatasetProfile]:
        tables, warnings, text = self._load(file_id, session_id)
        digest = self.store.get_file(file_id, session_id)['hash']
        result = []
        for name, frame in tables.items():
            missing = {str(k): int(v) for k, v in frame.isna().sum().items() if v}
            quality = [*warnings, f'Duplicate full rows: {int(frame.duplicated().sum())}', f'Missing cells by column: {json.dumps(missing, ensure_ascii=False)}']
            if frame.empty:
                quality.append('Table has no data rows')
            result.append(DatasetProfile(dataset_id=file_id, file_hash=digest, sheet_or_table=name,
                columns=list(frame.columns), dtypes={str(k): str(v) for k, v in frame.dtypes.items()},
                row_count=len(frame), quality_warnings=quality, preview=records(frame.head(5))))
        if not result:
            result.append(DatasetProfile(dataset_id=file_id, file_hash=digest, quality_warnings=[*warnings, f'No tables; text length={len(text)}']))
        return result

    def profile(self, file_id: str, session_id: str) -> DatasetProfile:
        profiles = self.profiles(file_id, session_id)
        first = profiles[0]
        if len(profiles) > 1:
            first.quality_warnings.append(f'Multiple tables: {[p.sheet_or_table for p in profiles]}; this profile describes only the first; use profiles/list_tables')
        return first

    def list_tables(self, file_id: str, session_id: str) -> list[dict]:
        return [p.model_dump() for p in self.profiles(file_id, session_id)]

    def read_text(self, file_id: str, session_id: str, max_chars: int = 12000) -> dict:
        if not 1 <= max_chars <= 50000:
            raise StoreError('max_chars must be between 1 and 50000', 422)
        _, _, text = self._load(file_id, session_id)
        return dict(file_id=file_id, text=text[:max_chars], total_characters=len(text), truncated=len(text) > max_chars)

    def read_result_table(self, artifact_id: str, session_id: str, limit: int = 30) -> dict:
        """Return explicitly labeled preview of a verified result artifact in this session."""
        artifact = self.store.get_artifact(artifact_id)
        if artifact.session_id != session_id or artifact.kind != 'table':
            raise StoreError('Table artifact not found in this session', 404, 'not_found')
        if not 1 <= limit <= 100:
            raise StoreError('Preview limit must be 1–100 rows', 422)
        path = self.store.resolve_artifact(artifact_id)
        key = (artifact_id, artifact.hash)
        with stage(self.store,'result_preview',artifact_id=artifact_id):
            if path.suffix.lower() == '.csv':
                frame, shape = self._csv_preview(path, limit, self._result_shapes.get(key))
            elif path.suffix.lower() == '.xlsx':
                frame, shape = self._xlsx_preview(path, limit, self._result_shapes.get(key))
            else:
                # Legacy XLS needs xlrd; retain its tested reader without
                # pretending pandas.read_excel has a chunksize parameter.
                frame = pd.read_excel(path, dtype=object, keep_default_na=False)
                shape = dict(columns=list(frame.columns), total_rows=len(frame))
                frame = infer_numbers(frame.replace('', None)).head(limit)
        self.store.resolve_artifact(artifact_id)
        self._result_shapes[key] = shape
        return dict(artifact_id=artifact_id, columns=list(frame.columns), total_rows=shape['total_rows'],
                    preview=records(frame), truncated=shape['total_rows'] > limit,
                    scope='preview; artifact contains full result')

    @staticmethod
    def _numeric_shape(frame, flags, floating, ranges):
        """Keep full-column inference, even when only the first rows are shown."""
        for column in frame.columns:
            if not flags.get(column, True):
                continue
            present = frame[column].dropna().astype(str)
            if present.empty:
                continue
            if present.str.match(r'^[+-]?0\d+').any():
                flags[column] = False
                continue
            try:
                numeric = pd.to_numeric(frame[column], errors='raise')
                if not pd.api.types.is_numeric_dtype(numeric.dtype):
                    flags[column] = False
                    continue
                floating[column] = floating.get(column, False) or pd.api.types.is_float_dtype(numeric.dtype)
                if pd.api.types.is_integer_dtype(numeric.dtype) and not numeric.empty:
                    old = ranges.get(column, (None, None))
                    low, high = int(numeric.min()), int(numeric.max())
                    ranges[column] = (min(old[0],low) if old[0] is not None else low,
                                      max(old[1],high) if old[1] is not None else high)
                flags[column] = True
            except (ValueError, TypeError):
                flags[column] = False

    @staticmethod
    def _preview_types(frame, shape):
        frame = frame.replace('', None)
        for column, numeric in shape.get('numeric', {}).items():
            low, high = shape.get('ranges', {}).get(column, (0,0))
            if low < 0 and high > 9223372036854775807 and not shape.get('floating', {}).get(column):
                # pandas keeps a mixed signed/uint64 column as strings; chunk
                # boundaries must not turn the first values into integers.
                numeric = False
            if numeric:
                frame[column] = pd.to_numeric(frame[column], errors='raise')
                if shape.get('floating', {}).get(column):
                    frame[column] = frame[column].astype(float)
        return frame

    @classmethod
    def _csv_preview(cls, path, limit, known=None):
        if known is not None:
            frame = pd.read_csv(path, encoding='utf-8-sig', dtype=str, keep_default_na=False, nrows=limit)
            return cls._preview_types(frame, known), known
        chunks, total, flags, floating, ranges, columns = [], 0, {}, {}, {}, None
        # pandas' documented iterator keeps memory bounded, including multiline
        # CSV records: https://pandas.pydata.org/docs/reference/api/pandas.read_csv.html
        with pd.read_csv(path, encoding='utf-8-sig', dtype=str, keep_default_na=False, chunksize=4096) as reader:
            for chunk in reader:
                columns = list(chunk.columns)
                chunk = chunk.replace('', None)
                cls._numeric_shape(chunk, flags, floating, ranges)
                if total < limit:
                    chunks.append(chunk.head(limit-total))
                total += len(chunk)
        if columns is None:
            frame = pd.read_csv(path, encoding='utf-8-sig', dtype=str, keep_default_na=False, nrows=0)
            columns = list(frame.columns)
        else:
            frame = pd.concat(chunks, ignore_index=True) if chunks else pd.DataFrame(columns=columns)
        shape = dict(columns=columns,total_rows=total,numeric=flags,floating=floating,ranges=ranges)
        return cls._preview_types(frame,shape), shape

    @classmethod
    def _xlsx_preview(cls, path, limit, known=None):
        if known is not None:
            from itertools import islice
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                columns = known['columns']
                rows = book.worksheets[0].iter_rows(min_row=2, values_only=True)
                captured = []
                for raw in islice(rows, min(limit, known['total_rows'])):
                    row = list(raw[:len(columns)])
                    captured.append(row + [None] * (len(columns)-len(row)))
                frame = pd.DataFrame(captured, columns=columns, dtype=object)
                return cls._preview_types(frame,known), known
            finally:
                book.close()
        book = load_workbook(path,read_only=True,data_only=True)
        try:
            sheet = book.worksheets[0]
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, ())
            # Use pandas' header interpretation on zero data rows (duplicate
            # labels and unnamed headings retain previous reader semantics).
            columns = list(pd.read_excel(path,dtype=object,keep_default_na=False,nrows=0).columns)
            captured, batch, total, last_nonempty, flags, floating, ranges = [], [], 0, 0, {}, {}, {}
            for raw in rows:
                row = list(raw[:len(columns)])
                row += [None] * (len(columns)-len(row))
                total += 1
                if any(value not in (None,'') for value in row):
                    last_nonempty = total
                if len(captured) < limit:
                    captured.append(row)
                batch.append(row)
                if len(batch) == 4096:
                    cls._numeric_shape(pd.DataFrame(batch,columns=columns,dtype=object).replace('',None),flags,floating,ranges)
                    batch.clear()
            if batch:
                cls._numeric_shape(pd.DataFrame(batch,columns=columns,dtype=object).replace('',None),flags,floating,ranges)
            frame = pd.DataFrame(captured[:min(limit,last_nonempty)],columns=columns,dtype=object)
            shape = dict(columns=columns,total_rows=last_nonempty,numeric=flags,floating=floating,ranges=ranges)
            return cls._preview_types(frame,shape),shape
        finally:
            book.close()

    @staticmethod
    def _columns(frame: pd.DataFrame, columns: list[str]):
        absent = sorted(set(columns) - set(frame.columns))
        if absent:
            raise StoreError(f'Unknown column(s) {absent}; available: {list(frame.columns)}', 422, 'unknown_column')

    def _filter(self, frame, filters):
        for condition in filters or []:
            if set(condition) - {'column', 'op', 'value'}:
                raise StoreError('Filter accepts column, op and value only', 422)
            column, op, value = condition.get('column'), condition.get('op'), condition.get('value')
            self._columns(frame, [column])
            series = frame[column]
            if op in {'eq', 'ne', 'gt', 'ge', 'lt', 'le'}:
                try:
                    mask = {'eq': series.eq, 'ne': series.ne, 'gt': series.gt, 'ge': series.ge, 'lt': series.lt, 'le': series.le}[op](value)
                except TypeError as exc:
                    raise StoreError('Filter value is incompatible with column type', 422) from exc
            elif op == 'in' and isinstance(value, list):
                mask = series.isin(value)
            elif op == 'contains':
                mask = series.astype('string').str.contains(str(value), regex=False, na=False)
            elif op == 'is_null':
                mask = series.isna()
            elif op == 'not_null':
                mask = series.notna()
            else:
                raise StoreError('Unsupported filter; use eq/ne/gt/ge/lt/le/in/contains/is_null/not_null', 422)
            frame = frame.loc[mask.fillna(False)]
        if frame.empty:
            raise StoreError('No rows match the supplied filters', 422, 'empty_result')
        return frame

    @traced_method('dataset_analysis')
    def analyze(self, task: TaskSpec, file_id: str, operation: str, table: str | None = None,
                columns: list[str] | None = None, filters: list[dict] | None = None,
                group_by: list[str] | None = None, aggregations: dict[str, Any] | None = None,
                sort_by: list[str] | None = None, ascending: bool = True,
                right_file_id: str | None = None, right_table: str | None = None,
                on: list[str] | None = None, how: str = 'inner', chart: str | None = None) -> AnalysisResult:
        if file_id not in task.file_ids or (right_file_id and right_file_id not in task.file_ids):
            raise StoreError('Only files selected for this task may be analyzed', 422, 'file_not_selected')
        source_ids = [file_id] + ([right_file_id] if right_file_id else [])
        frame = self._filter(self.read_table(file_id, task.session_id, table), filters)
        source_rows = len(frame)
        limitations = [w for p in self.profiles(file_id, task.session_id) for w in p.quality_warnings if 'formula' in w.lower()]
        metrics: dict = {'source_rows': source_rows}
        if operation == 'merge':
            if not right_file_id or not on or how not in {'inner', 'left', 'right', 'outer'}:
                raise StoreError('merge requires right_file_id, nonempty on keys, and inner/left/right/outer how', 422)
            right = self.read_table(right_file_id, task.session_id, right_table)
            self._columns(frame, on)
            self._columns(right, on)
            left_counts = frame.groupby(on, dropna=False).size().rename('left_n').reset_index()
            right_counts = right.groupby(on, dropna=False).size().rename('right_n').reset_index()
            key_counts = left_counts.merge(right_counts, on=on, how='outer').fillna(0)
            estimated = int((key_counts['left_n'] * key_counts['right_n']).sum())
            if how in {'left', 'outer'}:
                estimated += int(key_counts.loc[key_counts['right_n'].eq(0), 'left_n'].sum())
            if how in {'right', 'outer'}:
                estimated += int(key_counts.loc[key_counts['left_n'].eq(0), 'right_n'].sum())
            if estimated > 1000000 or estimated * (len(frame.columns) + len(right.columns) - len(on)) > 5000000:
                raise StoreError(f'Merge would create approximately {estimated} rows, exceeding resource limits; aggregate or deduplicate keys first', 422, 'merge_too_large')
            left_dup, right_dup = int(frame.duplicated(on).sum()), int(right.duplicated(on).sum())
            matched = frame.merge(right[on].drop_duplicates(), on=on, how='left', indicator=True)['_merge'].eq('both').mean()
            metrics.update(left_duplicate_keys=left_dup, right_duplicate_keys=right_dup, left_key_match_rate=float(matched), right_rows=len(right))
            frame = frame.merge(right, on=on, how=how)
            metrics['row_expansion_ratio'] = len(frame) / source_rows
            limitations.append(f'Merge duplicate key rows left={left_dup}, right={right_dup}; left match rate={matched:.6f}; output/source row ratio={len(frame)/source_rows:.6f}. Null keys match each other under pandas semantics.')
        elif operation == 'aggregate':
            if not aggregations:
                raise StoreError('aggregate requires aggregations, e.g. {"alcohol":"mean"}', 422)
            self._columns(frame, [*(group_by or []), *aggregations])
            allowed = {'sum', 'mean', 'count', 'min', 'max', 'median', 'nunique', 'std'}
            normalized = {k: [v] if isinstance(v, str) else v for k, v in aggregations.items()}
            for column, funcs in normalized.items():
                if not isinstance(funcs, list) or not funcs or any(f not in allowed for f in funcs):
                    raise StoreError(f'Allowed aggregation functions: {sorted(allowed)}', 422)
                if any(f in {'mean', 'median', 'sum', 'std'} for f in funcs) and not pd.api.types.is_numeric_dtype(frame[column]):
                    raise StoreError(f'Column {column!r} must be numeric for requested aggregation', 422)
            if group_by:
                frame = frame.groupby(group_by, dropna=False).agg(normalized)
                frame.columns = ['_'.join(c) for c in frame.columns]
                frame = frame.reset_index()
            else:
                values = {f'{col}_{func}': clean_json(getattr(frame[col], func)()) for col, funcs in normalized.items() for func in funcs}
                metrics.update(values)
                frame = pd.DataFrame([values])
        elif operation == 'describe':
            if columns:
                self._columns(frame, columns)
                frame = frame[columns]
            frame = frame.describe(include='all').rename_axis('statistic').reset_index()
        elif operation == 'correlation':
            selected = columns or list(frame.select_dtypes(include='number').columns)
            self._columns(frame, selected)
            if len(selected) < 2 or any(not pd.api.types.is_numeric_dtype(frame[c]) for c in selected):
                raise StoreError('correlation requires at least two numeric columns', 422)
            frame = frame[selected].corr().rename_axis('column').reset_index()
            limitations.append('Pearson pairwise-complete correlation; missing values excluded by pair; association does not establish causation')
        elif operation != 'rows':
            raise StoreError('operation must be rows/describe/aggregate/correlation/merge', 422)
        if frame.empty:
            raise StoreError('Operation produced no rows', 422, 'empty_result')
        if columns and operation in {'rows', 'merge'}:
            self._columns(frame, columns)
            frame = frame[columns]
        if sort_by:
            self._columns(frame, sort_by)
            frame = frame.sort_values(sort_by, ascending=ascending)
        directory = self.store.safe_path(f'script_output/{task.session_id}/{task.task_id}/tables')
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f'{operation}_{new_id()}.csv'
        frame.to_csv(path, index=False, encoding='utf-8-sig')
        artifact = self.store.register_artifact(path, task.task_id, task.session_id, 'table', source_ids)
        charts = []
        if chart:
            charts = [self._chart(task, frame, chart, source_ids)]
        metrics['result_rows'] = len(frame)
        definitions = {key: dict(operation=operation, table=table, filters=filters or [], group_by=group_by or [],
            aggregations=aggregations or {}, source_ids=source_ids, scope='full dataset; no sampling', missing='pandas: count/mean omit missing, group keys include missing') for key in metrics}
        result = AnalysisResult(task_id=task.task_id, session_id=task.session_id,
            input_versions={fid: self.store.get_file(fid, task.session_id)['hash'] for fid in source_ids},
            metrics=metrics, metric_definitions=definitions, table_ids=[artifact.artifact_id], chart_ids=charts,
            findings=[f'{operation}: {source_rows} source rows → {len(frame)} output rows; full result stored in table artifact'], limitations=limitations)
        for fid in source_ids:
            self.store.resolve_file(fid,task.session_id)
        return self.store.save_analysis(result)

    def _chart(self, task, frame, kind, source_ids):
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        from src.chart_export import configure_chart_fonts
        if kind not in {'bar', 'line', 'hist', 'scatter'}:
            raise StoreError('Chart kind must be bar/line/hist/scatter', 422)
        configure_chart_fonts()
        numeric = frame.select_dtypes(include='number').columns.tolist()
        if not numeric or (kind == 'scatter' and len(numeric) < 2):
            raise StoreError('Chart requires numeric columns (scatter requires two)', 422)
        fig, ax = plt.subplots(figsize=(8, 4.5))
        try:
            if kind == 'hist':
                ax.hist(frame[numeric[0]].dropna(), bins=20)
                ax.set_xlabel(numeric[0])
                ax.set_ylabel('Count')
                ax.set_title(f'{numeric[0]} 分布')
            elif kind == 'scatter':
                ax.scatter(frame[numeric[0]], frame[numeric[1]], s=12, alpha=.5)
                ax.set_xlabel(numeric[0]); ax.set_ylabel(numeric[1])
                ax.set_title(f'{numeric[0]} 与 {numeric[1]} 的关系')
            else:
                if len(frame) > 100:
                    raise StoreError('Aggregate or filter to at most 100 rows for bar/line charts; no implicit sampling', 422)
                x = frame.iloc[:, 0].astype(str) if len(frame.columns) > 1 else frame.index.astype(str)
                getattr(ax, 'plot' if kind == 'line' else kind)(x, frame[numeric[-1]])
                ax.set_xlabel(str(frame.columns[0]) if len(frame.columns) > 1 else 'Index')
                ax.set_ylabel(str(numeric[-1]))
                ax.set_title(str(numeric[-1]) + (' 变化趋势' if kind == 'line' else ' 对比'))
                ax.tick_params(axis='x', labelrotation=45)
            from src.chart_export import apply_report_theme, export_figure
            apply_report_theme(fig)
            fig.tight_layout()
            path = self.store.safe_path(f'script_output/{task.session_id}/{task.task_id}/charts/{new_id()}.png')
            path.parent.mkdir(parents=True, exist_ok=True)
            fig.savefig(path, dpi=140)
            from matplotlib.figure import Figure
            svg = path.with_suffix('.svg')
            export_figure(fig, svg, Figure.savefig)
            artifact = self.store.register_artifact(path, task.task_id, task.session_id, 'chart', source_ids)
            self.store.register_artifact(svg, task.task_id, task.session_id, 'chart_svg', [artifact.artifact_id])
            return artifact.artifact_id
        finally:
            plt.close(fig)
