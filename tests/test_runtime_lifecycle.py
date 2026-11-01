"""Independent report checks using real analysis/material tools and registered sources."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from src.agents import BudgetExceeded, TaskRuntime, analyst_tools, report_tools
from src.config import Settings
from src.data import DataService
from src.reporting import ReportService
from src.schemas import TaskSpec
from src.storage import Store, StoreError


@pytest.fixture
def context(tmp_path):
    store = Store(Settings(_env_file=None, workspace=tmp_path, api_key='review-offline-secret'))
    sid = store.new_session()['session_id']
    source = tmp_path / 'KB' / 'source.csv'
    source.write_text('id,value,text\n001,2,=SUM(A1)\nNA,3,<script>unsafe</script>\n', encoding='utf-8')
    fid = store.register_file(source, source.name, None, source='kb')['file_id']
    task = TaskSpec(session_id=sid, request='Analyze and report', file_ids=[fid], expected_outputs=['html', 'xlsx'])
    store.create_task(task)
    analysis = DataService(store).analyze(task, fid, 'aggregate', aggregations={'value': 'mean'})
    return store, task, fid, analysis


def tools(runtime):
    return {item.name: item for item in report_tools(runtime)}


@pytest.mark.parametrize('input_field,bad_text', [
    ('title', '均值为999'),
    ('summary', '均值为九百九十九。'),
    ('summary', '均值为９９９。'),
    ('summary', '均值为九。'),
])
def test_model_supplied_numeric_claim_requires_binding(context, input_field, bad_text):
    store, task, _, analysis = context
    kwargs = {'title': bad_text} if input_field == 'title' else {'narrative': {'summary': bad_text}}
    result = ReportService(store).generate(task, [analysis.analysis_id], **kwargs)
    assert result.status == 'need_analysis', (input_field, result)
    assert not result.report_ids
    assert not [a for a in store.list_artifacts(task.task_id) if a.kind.startswith('report')]


def test_qualitative_chinese_prose_and_bound_title_are_valid(context):
    store, task, _, analysis = context
    summary = '一致采用已确认口径，统一整理来源。一般先检查质量，一方面说明数据局限。'
    result = ReportService(store).generate(task, [analysis.analysis_id],
        title='均值为{{metric:' + analysis.analysis_id + ':value_mean}}', narrative={'summary': summary})
    assert result.status == 'completed', result
    html = next(store.resolve_artifact(aid) for aid in result.report_ids if store.resolve_artifact(aid).suffix == '.html').read_text(encoding='utf-8')
    soup = BeautifulSoup(html, 'html.parser')
    assert soup.h1.get_text() == '均值为2.5'
    assert summary in soup.get_text()


def test_real_tool_enforces_chosen_format_and_shared_supplement_limit(context):
    store, task, _, analysis = context
    runtime = TaskRuntime(task, store, lambda *a, **k: None, lambda: False)
    runtime.record_analysis(analysis)
    report = tools(runtime)
    mismatch = report['render_report'].invoke({'analysis_ids': [analysis.analysis_id], 'formats': ['html']})
    assert mismatch['code'] == 'report_format_mismatch'
    first = report['render_report'].invoke({'analysis_ids': []})
    second = report['render_report'].invoke({'analysis_ids': [analysis.analysis_id], 'narrative': {
        'sections': [{'metric_refs': [{'analysis_id': analysis.analysis_id, 'metric_key': 'missing_requested_metric'}]}]}})
    third = report['request_analysis'].invoke({'request': 'More data'})
    assert first['round'] == 1 and second['round'] == 2
    assert third['status'] == 'waiting_user' and runtime.clarification
    assert not runtime.reports


def test_default_template_disclosure_is_saved_inside_actual_report(context):
    store, task, _, analysis = context
    runtime = TaskRuntime(task, store, lambda *a, **k: None, lambda: False)
    runtime.record_analysis(analysis)
    render = tools(runtime)['render_report']
    bad = {'analysis_ids': [analysis.analysis_id], 'title': '均值999', 'narrative': {'summary': '均值999。'}}
    assert render.invoke(bad)['status'] == 'needs_revision'
    assert render.invoke(bad)['status'] == 'needs_revision'
    result = render.invoke({'analysis_ids': [analysis.analysis_id], 'narrative': None})
    assert result['status'] == 'completed', result
    warning = runtime.report_limitations[0]
    paths = {store.resolve_artifact(aid).suffix: store.resolve_artifact(aid) for aid in result['report_ids']}
    assert warning in paths['.html'].read_text(encoding='utf-8')
    workbook = load_workbook(paths['.xlsx'])
    try:
        assert warning in str(list(workbook['说明与口径'].values))
    finally:
        workbook.close()


def test_revision_cap_cannot_discard_a_missing_requested_metric(context):
    store, task, _, analysis = context
    runtime = TaskRuntime(task, store, lambda *a, **k: None, lambda: False)
    runtime.record_analysis(analysis)
    render = tools(runtime)['render_report']
    for _ in range(2):
        assert render.invoke({'analysis_ids': [analysis.analysis_id], 'narrative': {'summary': '均值999。'}})['status'] == 'needs_revision'
    missing = {'analysis_ids': [analysis.analysis_id], 'narrative': {'sections': [
        {'metric_refs': [{'analysis_id': analysis.analysis_id, 'metric_key': 'missing_requested_metric'}]}]}}
    result = render.invoke(missing)
    assert result['status'] == 'need_analysis', result
    assert runtime.report_rounds == 1
    assert not runtime.reports


def test_default_tool_budget_is_still_a_hard_shared_limit(context):
    store, task, _, _ = context
    store.settings.max_tool_calls = 32
    runtime = TaskRuntime(task, store, lambda *a, **k: None, lambda: False)
    for index in range(32):
        runtime.consume('main' if index % 2 else 'report_designer', 'review')
    with pytest.raises(BudgetExceeded):
        runtime.consume('analyst', 'review')
    assert runtime.calls == 32


def test_report_tool_rejects_same_session_unselected_analysis(context):
    store, task, fid, analysis = context
    report_task = TaskSpec(session_id=task.session_id, request='A separate report', file_ids=[fid])
    store.create_task(report_task)
    runtime = TaskRuntime(report_task, store, lambda *a, **k: None, lambda: False)
    result = tools(runtime)['render_report'].invoke({'analysis_ids': [analysis.analysis_id]})
    assert result['status'] == 'error'
    assert not runtime.reports


def test_html_and_excel_keep_actual_precision_lexemes_and_all_rows(context):
    store, task, fid, mean = context
    rows = DataService(store).analyze(task, fid, 'rows')
    result = ReportService(store).generate(task, [mean.analysis_id, rows.analysis_id],
        narrative={'summary': '实际均值{{metric:' + mean.analysis_id + ':value_mean}}。'},
        layout={'include_metrics_summary': True, 'include_tables': True})
    paths = {store.resolve_artifact(aid).suffix: store.resolve_artifact(aid) for aid in result.report_ids}
    html = paths['.html'].read_text(encoding='utf-8')
    soup = BeautifulSoup(html, 'html.parser')
    assert '实际均值2.5。' in soup.get_text()
    assert not soup.find('script') and '&lt;script&gt;' in html
    book = load_workbook(paths['.xlsx'], data_only=False)
    try:
        sheet = book['结果表_2_1']
        assert sheet.max_row == 3
        assert sheet['A2'].value == '001' and sheet['A3'].value == 'NA'
        assert sheet['C2'].value == '=SUM(A1)' and sheet['C2'].data_type == 's'
        assert all(cell.data_type != 'f' for page in book for row in page for cell in row)
        exact = next(row[3] for row in book['指标汇总'].values if row[1] == 'value_mean')
        assert exact == '2.5'
    finally:
        book.close()


@pytest.mark.parametrize('suffix', ['.html', '.md', '.txt'])
def test_real_prepare_material_tool_retains_source_and_nonstatistical_scope(context, suffix):
    store, task, _, _ = context
    text = '材料记载收入为999元；这是一项待核实的来源陈述。' * 350
    source = store.root / 'KB' / ('brief' + suffix)
    source.write_text('<p>' + text + '</p>' if suffix == '.html' else text, encoding='utf-8')
    fid = store.register_file(source, source.name, None, source='kb')['file_id']
    task.file_ids = [fid]
    store.update_task(task.task_id, spec=task.model_dump())
    runtime = TaskRuntime(task, store, lambda *a, **k: None, lambda: False)
    prepare = next(item for item in analyst_tools(runtime) if item.name == 'prepare_material')
    prepared = prepare.invoke({'file_id': fid})
    assert 'analysis_id' in prepared, prepared
    analysis = store.get_analysis(prepared['analysis_id'], task.session_id)
    assert analysis.input_versions == {fid: store.get_file(fid, task.session_id)['hash']}
    assert analysis.metric_definitions['source_characters']['operation'] == 'prepare_material'
    result = tools(runtime)['render_report'].invoke({'analysis_ids': [analysis.analysis_id]})
    assert result['status'] == 'completed', result
    paths = [store.resolve_artifact(aid) for aid in result['report_ids']]
    html = next(path for path in paths if path.suffix == '.html').read_text(encoding='utf-8')
    assert '原文已截断' in html and '未进行业务统计' in html
    assert '不进行数据统计或验证原文主张' in html
    manifest = json.loads(store.resolve_artifact(result['validation_summary']['manifest_id']).read_text(encoding='utf-8'))
    assert manifest['sources'][0]['file_id'] == fid


def test_failed_new_version_preserves_prior_published_reports(context, monkeypatch):
    store, task, _, analysis = context
    first = ReportService(store).generate(task, [analysis.analysis_id])
    previous = {aid: store.resolve_artifact(aid).read_bytes() for aid in first.report_ids}
    index_before = set(store._index('artifacts'))
    register = store.register_artifact

    def fail_xlsx(path, *args, **kwargs):
        if Path(path).suffix == '.xlsx':
            raise RuntimeError('simulated second output registration failure')
        return register(path, *args, **kwargs)

    monkeypatch.setattr(store, 'register_artifact', fail_xlsx)
    with pytest.raises(RuntimeError, match='second output'):
        ReportService(store).generate(task, [analysis.analysis_id])
    assert set(store._index('artifacts')) == index_before
    assert all(store.resolve_artifact(aid).read_bytes() == content for aid, content in previous.items())
    assert not list((store.root / 'script_output').rglob('.report_staging_*'))


def test_output_source_changed_during_render_cannot_be_published(context, monkeypatch):
    store, task, _, analysis = context
    service = ReportService(store)
    validate = service._validate
    source_table = store.resolve_artifact(analysis.table_ids[0])

    def change_output(*args):
        validate(*args)
        source_table.write_text('value_mean\n999\n', encoding='utf-8')

    monkeypatch.setattr(service, '_validate', change_output)
    with pytest.raises(StoreError):
        service.generate(task, [analysis.analysis_id])
    assert not [item for item in store.list_artifacts(task.task_id) if item.kind.startswith('report')]
    assert not list((store.root / 'report').rglob('report.*'))
