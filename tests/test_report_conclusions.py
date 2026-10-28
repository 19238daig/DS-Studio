"""Reports answer the question without automatically embedding result CSVs."""
import csv
import json
from decimal import Decimal

import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from src.config import Settings
from src.data import DataService
from src.reporting import ReportService
from src.schemas import TaskSpec
from src.storage import Store, StoreError, file_hash


@pytest.fixture
def relationship(tmp_path):
    store = Store(Settings(workspace=tmp_path, _env_file=None))
    sid = store.new_session()['session_id']
    path = tmp_path / 'KB' / 'customers.csv'
    path.write_text('loyalty_tier_final,customer,revenue\ngold,A,120\ngold,B,80\nsilver,C,50\nsilver,D,30\n', encoding='utf-8')
    fid = store.refresh_kb()[0]['file_id']
    task = TaskSpec(session_id=sid, request='生成客户与销售之间的关系报告', file_ids=[fid])
    store.create_task(task)
    analysis = DataService(store).analyze(task, fid, 'aggregate',
        group_by=['loyalty_tier_final'], aggregations={'revenue': 'sum'}, chart='bar')
    # Compute from the complete immutable source, not a preview or model claim.
    with path.open(encoding='utf-8', newline='') as stream:
        rows = list(csv.DictReader(stream))
    totals = {tier: sum((Decimal(row['revenue']) for row in rows if row['loyalty_tier_final'] == tier), Decimal(0))
              for tier in ('gold', 'silver')}
    analysis.metrics.update(gold_revenue=str(totals['gold']), silver_revenue=str(totals['silver']))
    analysis.metric_definitions.update({f'{tier}_revenue': {'label': f'{tier}客户销售额', 'definition': '完整数据按等级汇总revenue'}
                                        for tier in totals})
    analysis.findings = [f"gold客户销售额为{totals['gold']}，高于silver客户的{totals['silver']}；该差异不证明会员等级导致销售变化。"]
    store.save_analysis(analysis)
    return store, task, analysis


def paths(store, result):
    return {store.resolve_artifact(aid).suffix: store.resolve_artifact(aid) for aid in result.report_ids}


def test_new_report_starts_with_actual_conclusion_without_group_csv(relationship):
    store, task, analysis = relationship
    result = ReportService(store).generate(task, [analysis.analysis_id])
    reports = paths(store, result)
    soup = BeautifulSoup(reports['.html'].read_text(encoding='utf-8'), 'html.parser')
    assert soup.find('h2').get_text() == '结论'
    conclusion = soup.select_one('.report-conclusion').get_text()
    assert 'gold客户销售额为200' in conclusion and '高于silver客户销售额的80' in conclusion
    assert not soup.find('h2', string='结果明细') and not soup.find('h2', string='指标汇总')
    assert not soup.find('table')  # Chart labels still describe the plotted groups.
    assert store.get_artifact(analysis.table_ids[0]).name not in soup.get_text()
    assert len(soup.select('svg[data-ds-chart]')) == 1
    assert soup.select_one('#ds-chart-interactions')
    book = load_workbook(reports['.xlsx'])
    try:
        assert not any(name.startswith('结果表_') for name in book.sheetnames)
        assert dict(book['说明与口径'].values)['结论'] == conclusion
    finally:
        book.close()
    manifest = json.loads(store.resolve_artifact(result.validation_summary['manifest_id']).read_text(encoding='utf-8'))
    assert manifest['analyses'][0]['table_ids'] == analysis.table_ids
    assert manifest['tables'] and manifest['metric_bindings']
    assert manifest['validation']['result_details_included'] is False
    assert store.resolve_artifact(analysis.table_ids[0]).exists()


@pytest.mark.parametrize('narrative', [None, {'summary': ''}, {'summary': '本报告基于已验证分析结果整理，指标与来源可逐项追溯。'}])
def test_safe_fallback_uses_business_metrics_not_execution_metadata(relationship, narrative):
    store, task, analysis = relationship
    analysis.findings = ['aggregate: 4 source rows → 2 output rows; full result stored in table artifact']
    store.save_analysis(analysis)
    result = ReportService(store).generate(task, [analysis.analysis_id], narrative=narrative)
    soup = BeautifulSoup(paths(store, result)['.html'].read_text(encoding='utf-8'), 'html.parser')
    conclusion = soup.select_one('.report-conclusion').get_text()
    assert 'gold客户销售额为200' in conclusion and 'silver客户销售额的80' in conclusion
    assert 'source rows' not in conclusion and 'source_rows' not in conclusion
    assert '本报告基于' not in conclusion


def test_explicit_details_and_revisions_preserve_original_and_layout(relationship):
    store, task, analysis = relationship
    service = ReportService(store)
    original = service.generate(task, [analysis.analysis_id], layout={'include_tables': True})
    original_paths = paths(store, original)
    before = {suffix: file_hash(path) for suffix, path in original_paths.items()}
    store.update_task(task.task_id, status='succeeded', result=original.model_dump())
    edit = TaskSpec(session_id=task.session_id, request='只修改标题', analysis_ids=[analysis.analysis_id])
    store.create_task(edit)
    revised = service.generate(edit, [analysis.analysis_id], title='客户销售结论', base_report_id=original.report_ids[0])
    assert revised.validation_summary['layout']['include_tables'] is True
    assert '结果明细' in paths(store, revised)['.html'].read_text(encoding='utf-8')
    removed = service.generate(edit, [analysis.analysis_id], base_report_id=original.report_ids[0], layout={'include_tables': False})
    assert '结果明细' not in paths(store, removed)['.html'].read_text(encoding='utf-8')
    assert before == {suffix: file_hash(path) for suffix, path in original_paths.items()}
    assert set(original.report_ids).isdisjoint(revised.report_ids + removed.report_ids)


def test_unbound_summary_is_still_rejected_without_report_files(relationship):
    store, task, analysis = relationship
    result = ReportService(store).generate(task, [analysis.analysis_id], narrative={'summary': 'gold客户收入为999。'})
    assert result.status == 'need_analysis' and not result.report_ids
    assert result.validation_summary['reason_code'] == 'unbound_prose'


def test_hidden_tables_keep_provenance_without_workbook_export_limits(relationship):
    store, task, analysis = relationship
    path = store.safe_path(f'script_output/{task.session_id}/{task.task_id}/wide.csv')
    width = 16385  # Wider than Excel; the default report does not export it.
    path.write_text(','.join(f'c{i}' for i in range(width)) + '\n' + ','.join('0' for _ in range(width)) + '\n', encoding='utf-8')
    artifact = store.register_artifact(path, task.task_id, task.session_id, 'table', list(analysis.input_versions))
    analysis.table_ids = [artifact.artifact_id]
    store.save_analysis(analysis)
    service = ReportService(store)
    result = service.generate(task, [analysis.analysis_id])
    assert result.status == 'completed'
    manifest = json.loads(store.resolve_artifact(result.validation_summary['manifest_id']).read_text(encoding='utf-8'))
    assert manifest['tables'][0]['total_rows'] == 1
    assert manifest['tables'][0]['shown_rows'] == 0
    assert manifest['tables'][0]['hash'] == artifact.hash
    with pytest.raises(StoreError, match='export cell limits'):
        service.generate(task, [analysis.analysis_id], layout={'include_tables': True})


def test_comparison_fallback_prioritizes_requested_measure_with_exact_values(relationship):
    store, task, analysis = relationship
    task.request = '比较甲乙方案的平均响应时间，直接给出高低结论'
    analysis.metrics = {'total_requests': 1000, 'success_count': 990,
                        'alpha_mean_latency': '18.1234567890123456789',
                        'beta_mean_latency': '12.9876543210987654321'}
    analysis.metric_definitions = {
        'alpha_mean_latency': {'label': '甲方案平均响应时间', 'unit': 'ms'},
        'beta_mean_latency': {'label': '乙方案平均响应时间', 'unit': 'ms'}}
    analysis.findings = ['total_requests为1000。', 'success_count为990。',
                         '甲方案耗时约18.12，是乙方案的1.4倍，金额为1000 USD。']
    store.save_analysis(analysis)
    result = ReportService(store).generate(task, [analysis.analysis_id])
    soup = BeautifulSoup(paths(store, result)['.html'].read_text(encoding='utf-8'), 'html.parser')
    conclusion = soup.select_one('.report-conclusion').get_text()
    assert '甲方案平均响应时间为18.1234567890123456789' in conclusion
    assert '高于乙方案平均响应时间的12.9876543210987654321' in conclusion
    assert 'total_requests' not in conclusion and '1000' not in conclusion
    assert '1.4倍' not in conclusion and 'USD' not in conclusion


def test_fallback_does_not_invent_comparison_when_only_counts_exist(relationship):
    store, task, analysis = relationship
    task.request = '比较甲乙方案平均响应时间'
    analysis.metrics = {'total_requests': 1000}
    analysis.metric_definitions = {'total_requests': {'label': '请求总数'}}
    analysis.findings = ['total_requests为1000。', '甲比乙快2倍，平均耗时12.34 ms。']
    store.save_analysis(analysis)
    result = ReportService(store).generate(task, [analysis.analysis_id])
    conclusion = BeautifulSoup(paths(store, result)['.html'].read_text(encoding='utf-8'), 'html.parser').select_one('.report-conclusion').get_text()
    assert '2倍' not in conclusion and '12.34' not in conclusion
    assert '高于' not in conclusion and '低于' not in conclusion
    assert '不足以完成' in conclusion


def test_fallback_labels_cannot_add_unbound_numbers_and_mixed_units_are_not_compared(relationship):
    store, task, analysis = relationship
    task.request = '比较甲乙平均响应时间'
    analysis.metrics = {'alpha_mean_latency': '18.5', 'beta_mean_latency': '12.5'}
    analysis.metric_definitions = {'alpha_mean_latency': {'label': '甲方案999倍响应', 'unit': 'ms'},
                                   'beta_mean_latency': {'label': '乙方案响应时间', 'unit': 'seconds'}}
    analysis.findings = []
    store.save_analysis(analysis)
    result = ReportService(store).generate(task, [analysis.analysis_id])
    conclusion = BeautifulSoup(paths(store, result)['.html'].read_text(encoding='utf-8'), 'html.parser').select_one('.report-conclusion').get_text()
    assert '999' not in conclusion and '高于' not in conclusion and '低于' not in conclusion


def test_shared_average_operator_does_not_compare_different_quantities(relationship):
    store, task, analysis = relationship
    task.request = '比较平均支付金额与平均响应时间'
    analysis.metrics = {'avg_payment': '200', 'avg_latency': '18.5'}
    analysis.metric_definitions = {'avg_payment': {'label': '平均支付金额'},
                                   'avg_latency': {'label': '平均响应时间'}}
    analysis.findings = ['该方案高一百倍。']
    store.save_analysis(analysis)
    result = ReportService(store).generate(task, [analysis.analysis_id])
    conclusion = BeautifulSoup(paths(store, result)['.html'].read_text(encoding='utf-8'), 'html.parser').select_one('.report-conclusion').get_text()
    assert '高于' not in conclusion and '低于' not in conclusion and '一百倍' not in conclusion
    assert '不足以完成' in conclusion
