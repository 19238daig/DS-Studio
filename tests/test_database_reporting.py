"""Report DB provenance with real readonly PostgreSQL and temporary artifacts."""
import json

import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from src.config import Settings
from src.data import DataService
from src.database import DatabaseService
from src.reporting import ReportService
from src.schemas import AnalysisResult, TaskSpec
from src.storage import Store, StoreError, file_hash


@pytest.fixture
def database_report(tmp_path):
    # Reads the existing local readonly profile; never copies credentials into
    # generated artifacts, pytest parametrization, or test output.
    settings = Settings(workspace=tmp_path)
    if not settings.db_enabled:
        pytest.skip('Local PostgreSQL source is not configured')
    store = Store(settings)
    sid = store.new_session()['session_id']
    task = TaskSpec(session_id=sid, request='客户订单与付款报告',
        database_source_ids=['olist'], database_tables={'olist':['orders','order_payments']})
    store.create_task(task)
    results = DatabaseService(store).query_batch(task, 'olist', [
        {'sql':'SELECT COUNT(*) AS order_count FROM olist.orders'},
        {'sql':'SELECT SUM(payment_value) AS payment_total FROM olist.order_payments'},
    ])
    metrics = {key:value for result in results for key,value in result['scalar_metrics'].items()}
    analysis = AnalysisResult(task_id=task.task_id, session_id=sid,
        db_refs=[r['db_ref'] for r in results], table_ids=[r['result_artifact_id'] for r in results],
        metrics=metrics, metric_definitions={key:{'source':'registered SQL snapshot'} for key in metrics},
        findings=[f"订单数为{metrics['order_count']}，付款总额为{metrics['payment_total']}。"])
    store.save_analysis(analysis)
    return store, task, analysis, results


def deliver(store, task, result):
    store.update_task(task.task_id, status='succeeded', result=result.model_dump())


def report_files(store, result):
    return {store.get_artifact(aid).name:store.resolve_artifact(aid) for aid in result.report_ids}


def test_db_only_report_contains_actual_sql_scope_hashes_in_html_excel_manifest(database_report):
    store, task, analysis, results = database_report
    result = ReportService(store).generate(task, [analysis.analysis_id])
    assert result.status == 'completed'
    files = report_files(store, result)
    html = files['report.html'].read_text(encoding='utf-8')
    soup = BeautifulSoup(html, 'html.parser')
    assert soup.find('h2').get_text() == '结论'
    assert not soup.find('h2', string='指标汇总') and not soup.find('h2', string='结果明细')
    assert 'file_id:' not in html and 'PostgreSQL' in html
    assert str(analysis.metrics['order_count']) in soup.select_one('.report-conclusion').get_text()
    manifest = json.loads(store.resolve_artifact(result.validation_summary['manifest_id']).read_text(encoding='utf-8'))
    assert manifest['analyses'][0]['input_versions'] == {}
    assert manifest['analyses'][0]['db_refs'] == analysis.db_refs
    assert len(manifest['sources']) == 2
    expected_artifacts = {r[k] for r in results for k in ('result_artifact_id','sql_artifact_id','metadata_artifact_id')}
    assert expected_artifacts.issubset({a['artifact_id'] for a in manifest['source_artifacts']})
    for snapshot, source in zip(results, manifest['sources']):
        assert 'file_id' not in source
        assert source['query_id'] == snapshot['query_id']
        assert source['result_hash'] == snapshot['result_hash']
        assert source['sql_hash'] == snapshot['sql_hash']
        assert source['metadata_hash'] == store.get_artifact(snapshot['metadata_artifact_id']).hash
        assert source['queried_at'] == snapshot['queried_at']
        assert source['scope_fingerprint'] == snapshot['scope_fingerprint']
        assert snapshot['sql'] in soup.get_text()
        assert snapshot['queried_at'] in html and snapshot['result_hash'] in html
    book = load_workbook(files['report.xlsx'], read_only=True)
    try:
        assert '指标汇总' not in book.sheetnames
        assert not any(name.startswith('结果表_') for name in book.sheetnames)
        rows = list(book['数据库查询'].values)
        assert len(rows) == 3
        assert rows[1][1] == results[0]['query_id']
        assert rows[2][7] == results[1]['sql']
        assert rows[2][9] == results[1]['result_hash']
    finally:
        book.close()


def test_db_report_revision_keeps_originals_and_never_reruns_business_sql(database_report, monkeypatch):
    store, task, analysis, _ = database_report
    service = ReportService(store)
    original = service.generate(task, [analysis.analysis_id])
    deliver(store, task, original)
    old_files = report_files(store, original)
    old_hashes = {name:file_hash(path) for name,path in old_files.items()}
    registered_queries = set(store._index('database_queries'))
    revise = TaskSpec(session_id=task.session_id, request='只改标题', analysis_ids=[analysis.analysis_id],
        database_source_ids=task.database_source_ids, database_tables=task.database_tables)
    store.create_task(revise)
    def no_business_query(*args, **kwargs):
        raise AssertionError('Report revisions must reuse immutable SQL results')
    monkeypatch.setattr(DatabaseService, 'query', no_business_query)
    monkeypatch.setattr(DatabaseService, 'query_batch', no_business_query)
    revised = service.generate(revise, [analysis.analysis_id], title='客户订单管理简报',
        base_report_id=original.report_ids[0], layout={'include_quality_and_limitations':False})
    assert revised.status == 'completed'
    assert set(store._index('database_queries')) == registered_queries
    assert old_hashes == {name:file_hash(path) for name,path in old_files.items()}
    assert not set(original.report_ids) & set(revised.report_ids)
    assert '客户订单管理简报' in report_files(store, revised)['report.html'].read_text(encoding='utf-8')
    manifest = json.loads(store.resolve_artifact(revised.validation_summary['manifest_id']).read_text(encoding='utf-8'))
    assert manifest['revision_of']['hash'] == old_hashes['report.html']
    assert manifest['analyses'][0]['db_refs'] == analysis.db_refs


@pytest.mark.parametrize('selection', ['restricted_source', 'narrower_tables', 'no_sources'])
def test_report_current_scope_cannot_reuse_full_snapshot(database_report, selection):
    store, task, analysis, _ = database_report
    store.update_task(task.task_id, status='succeeded', result={'analysis_ids':[analysis.analysis_id]})
    selected = ['sellers_restricted'] if selection == 'restricted_source' else (['olist'] if selection == 'narrower_tables' else [])
    report_task = TaskSpec(session_id=task.session_id, request='受限报告', analysis_ids=[analysis.analysis_id],
        database_source_ids=selected, database_tables={'olist':['orders']} if selection == 'narrower_tables' else {})
    store.create_task(report_task)
    with pytest.raises(StoreError) as failure:
        ReportService(store).generate(report_task, [analysis.analysis_id])
    assert failure.value.code in {'database_source_not_selected','database_table_not_allowed'}


@pytest.mark.parametrize('artifact_key', ['result_artifact_id','sql_artifact_id','metadata_artifact_id'])
def test_tampered_database_snapshot_artifacts_cannot_be_reported(database_report, artifact_key):
    store, task, analysis, results = database_report
    path = store.resolve_artifact(results[0][artifact_key])
    with path.open('ab') as stream:
        stream.write(b'\nTAMPERED')
    with pytest.raises(StoreError):
        ReportService(store).generate(task, [analysis.analysis_id])
    assert not list((store.root/'report').rglob('report.*'))


def file_report_fixture(tmp_path):
    store = Store(Settings(workspace=tmp_path, _env_file=None))
    sid = store.new_session()['session_id']
    path = store.root/'KB'/'legacy.csv'
    path.write_text('amount\n10\n20\n', encoding='utf-8')
    fid = store.refresh_kb()[0]['file_id']
    task = TaskSpec(session_id=sid, request='文件报告', file_ids=[fid])
    store.create_task(task)
    analysis = DataService(store).analyze(task, fid, 'aggregate', aggregations={'amount':'mean'})
    service = ReportService(store)
    report = service.generate(task, [analysis.analysis_id])
    deliver(store, task, report)
    return store, task, analysis, report, service


@pytest.mark.parametrize('tamper', [False, True])
def test_historical_manifest_normalizes_added_defaults_without_accepting_tampering(tmp_path, tamper):
    store, task, analysis, report, service = file_report_fixture(tmp_path)
    artifact_id = report.validation_summary['manifest_id']
    path = store.resolve_artifact(artifact_id)
    manifest = json.loads(path.read_text(encoding='utf-8'))
    manifest['analyses'][0].pop('db_refs')  # Exact shape stored before DB support.
    if tamper:
        manifest['analyses'][0]['metrics']['amount_mean'] = 999
    path.write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
    # Model a legitimately registered pre-upgrade manifest, with its actual
    # immutable hash. read_report must still compare every analysis field.
    index = store._index('artifacts')
    index[artifact_id].update(hash=file_hash(path), size=path.stat().st_size)
    store._save_index('artifacts', index)
    if tamper:
        with pytest.raises(StoreError) as failure:
            service.read_report(task, report.report_ids[0])
        assert failure.value.code == 'stale_report_analysis'
    else:
        assert service.read_report(task, report.report_ids[0])['analysis_ids'] == [analysis.analysis_id]


def test_file_only_reports_do_not_connect_to_database(tmp_path, monkeypatch):
    def unavailable(*args, **kwargs):
        raise AssertionError('File-only report attempted to access PostgreSQL')
    monkeypatch.setattr(DatabaseService, 'authorize_snapshot', unavailable)
    store, task, _, report, service = file_report_fixture(tmp_path)
    assert report.status == 'completed'
    assert service.read_report(task, report.report_ids[0])['analysis_ids']
