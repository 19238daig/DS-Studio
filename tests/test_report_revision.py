import json

import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from src.config import Settings
from src.data import DataService
from src.reporting import ReportService
from src.schemas import TaskSpec
from src.storage import Store, StoreError, file_hash


@pytest.fixture
def reports(tmp_path):
    store = Store(Settings(workspace=tmp_path, _env_file=None))
    sid = store.new_session()['session_id']
    path = tmp_path / 'KB' / 'values.csv'
    path.write_text('amount,other\n10,2\n20,3\n30,5\n', encoding='utf-8')
    fid = store.refresh_kb()[0]['file_id']
    original = TaskSpec(session_id=sid, request='生成报告', file_ids=[fid])
    store.create_task(original)
    data = DataService(store)
    a = data.analyze(original, fid, 'aggregate', aggregations={'amount': 'mean'})
    b = data.analyze(original, fid, 'correlation', columns=['amount', 'other'], chart='bar')
    ids = [a.analysis_id, b.analysis_id]
    narrative = {'summary': '均值为{{metric:' + a.analysis_id + ':amount_mean}}。',
                 'sections': [{'heading': '保留章节', 'text': '完整数据计算结果。',
                               'metric_refs': [{'analysis_id': a.analysis_id, 'metric_key': 'amount_mean', 'label': '均值'}]}]}
    service = ReportService(store)
    # A historical report with an explicitly visible metrics module.
    source = service.generate(original, ids, title='原始报告', narrative=narrative,
                              layout={'include_metrics_summary': True, 'include_tables': True})
    store.update_task(original.task_id, status='succeeded', result={'analysis_ids': ids, 'report_ids': source.report_ids})
    task = TaskSpec(session_id=sid, request='删除指标汇总', analysis_ids=ids)
    store.create_task(task)
    return store, service, original, source, task


def deliver(store, task, result):
    store.update_task(task.task_id, status='succeeded', result=result.model_dump())


def test_remove_metrics_really_changes_both_formats_and_preserves_everything_else(reports):
    store, service, _, source, task = reports
    old_paths = [store.resolve_artifact(rid) for rid in source.report_ids]
    old_hashes = [file_hash(path) for path in old_paths]
    base = service.read_report(task, source.report_ids[0])
    result = service.generate(task, base['analysis_ids'], base_report_id=base['report_id'], layout={'include_metrics_summary': False})
    assert result.status == 'completed' and len(result.report_ids) == 2
    assert result.validation_summary['layout_changes'] == {'include_metrics_summary': False}
    html = next(store.resolve_artifact(rid) for rid in result.report_ids if store.get_artifact(rid).name.endswith('.html'))
    soup = BeautifulSoup(html.read_text(encoding='utf-8'), 'html.parser')
    assert not soup.find('h2', string='指标汇总') and not soup.select('.metric-value')
    assert soup.h1.get_text() == '原始报告' and '均值为20.0。' in soup.get_text()
    assert soup.find('h2', string='保留章节') and len(soup.select('svg[data-ds-chart]')) == 1
    xlsx = next(store.resolve_artifact(rid) for rid in result.report_ids if store.get_artifact(rid).name.endswith('.xlsx'))
    book = load_workbook(xlsx)
    try:
        assert '指标汇总' not in book.sheetnames
        assert book['结果表_1_1'].max_row == 2 and '章节_1' in book.sheetnames
        assert len(book['图表']._images) == 1
    finally:
        book.close()
    manifest = json.loads(store.resolve_artifact(result.validation_summary['manifest_id']).read_text(encoding='utf-8'))
    assert manifest['revision_of']['report_id'] == base['report_id']
    assert [a['analysis_id'] for a in manifest['analyses']] == base['analysis_ids']
    assert manifest['metric_bindings'] and not manifest['validation']['metrics_summary_included']
    assert old_hashes == [file_hash(path) for path in old_paths]


def test_later_title_edit_inherits_removal_and_restore_is_explicit(reports):
    store, service, _, source, task = reports
    base = service.read_report(task, source.report_ids[0])
    first = service.generate(task, base['analysis_ids'], base_report_id=base['report_id'], layout={'include_metrics_summary': False})
    deliver(store, task, first)
    next_task = TaskSpec(session_id=task.session_id, request='只改标题')
    store.create_task(next_task)
    second = service.generate(next_task, base['analysis_ids'], title='管理层简报', base_report_id=first.report_ids[0])
    html = store.resolve_artifact(second.report_ids[0]).read_text(encoding='utf-8')
    assert '管理层简报' in html and '<h2 data-section="metrics_summary">' not in html
    deliver(store, next_task, second)
    restore = TaskSpec(session_id=task.session_id, request='恢复指标汇总')
    store.create_task(restore)
    third = service.generate(restore, base['analysis_ids'], base_report_id=second.report_ids[0], layout={'include_metrics_summary': True})
    assert 'metrics_summary' in store.resolve_artifact(third.report_ids[0]).read_text(encoding='utf-8')


@pytest.mark.parametrize('issue', ['other_session', 'tampered_report', 'analysis_subset', 'wrong_kind'])
def test_revision_base_does_not_bypass_provenance(reports, issue):
    store, service, _, source, task = reports
    rid = source.report_ids[0]
    if issue == 'other_session':
        task = TaskSpec(session_id=store.new_session()['session_id'], request='edit')
        store.create_task(task)
    elif issue == 'tampered_report':
        store.resolve_artifact(rid).write_text('tampered', encoding='utf-8')
    elif issue == 'wrong_kind':
        rid = store.get_analysis(task.analysis_ids[0], task.session_id).table_ids[0]
    if issue == 'analysis_subset':
        with pytest.raises(StoreError, match='preserve all'):
            service.generate(task, task.analysis_ids[:1], base_report_id=rid)
    else:
        with pytest.raises(StoreError):
            service.read_report(task, rid)


@pytest.mark.parametrize('formats', [['html'], ['xlsx']])
def test_legacy_report_without_presentation_snapshot_can_be_revised(reports, monkeypatch, formats):
    store, service, original, _, task = reports
    register = store.register_artifact
    def legacy_register(path, *args, **kwargs):
        if path.name == 'manifest.json':
            value = json.loads(path.read_text(encoding='utf-8'))
            value.pop('presentation')
            path.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
        return register(path, *args, **kwargs)
    monkeypatch.setattr(store, 'register_artifact', legacy_register)
    first = service.generate(original, task.analysis_ids, title='旧版报告', narrative={'summary': '已有可信摘要。'}, formats=formats,
                             layout={'include_metrics_summary': True})
    deliver(store, original, first)
    monkeypatch.setattr(store, 'register_artifact', register)
    base = service.read_report(task, first.report_ids[0])
    assert base['presentation']['summary'] == '已有可信摘要。'
    result = service.generate(task, task.analysis_ids, base_report_id=base['report_id'], layout={'include_metrics_summary': False})
    assert len(result.report_ids) == 1 and result.validation_summary['formats'] == formats
    assert result.validation_summary['layout'] == {**base['presentation']['layout'], 'include_metrics_summary': False}


def test_invalid_layout_never_silently_defaults_to_original(reports):
    _, service, _, source, task = reports
    with pytest.raises(StoreError, match='layout'):
        service.generate(task, task.analysis_ids, base_report_id=source.report_ids[0], layout={'include_metrics_summary': 'false'})


def test_quality_limitations_and_sources_are_removed_from_both_formats_but_retained_in_manifest(reports):
    store,service,_,source,task=reports
    old_hashes=[file_hash(store.resolve_artifact(rid)) for rid in source.report_ids]
    result=service.generate(task,task.analysis_ids,base_report_id=source.report_ids[0],
        layout={'include_quality_and_limitations':False,'include_sources':False,'include_metrics_summary':True})
    html=store.resolve_artifact(result.report_ids[0]).read_text(encoding='utf-8')
    soup=BeautifulSoup(html,'html.parser')
    assert not soup.find('h2',string='数据质量与局限') and not soup.find('h2',string='来源与可追溯信息')
    assert soup.find('h2',string='指标汇总') and len(soup.select('svg[data-ds-chart]'))==1
    book=load_workbook(store.resolve_artifact(result.report_ids[1]),read_only=True)
    try:
        assert '数据质量' not in book.sheetnames and '数据来源' not in book.sheetnames
        assert not any(row[0]=='局限' for row in book['说明与口径'].values)
        assert '指标汇总' in book.sheetnames and '结果表_1_1' in book.sheetnames
    finally: book.close()
    manifest=json.loads(store.resolve_artifact(result.validation_summary['manifest_id']).read_text(encoding='utf-8'))
    assert manifest['sources'] and manifest['limitations'] and manifest['metric_bindings']
    assert old_hashes==[file_hash(store.resolve_artifact(rid)) for rid in source.report_ids]
    deliver(store,task,result)
    next_task=TaskSpec(session_id=task.session_id,request='只改标题')
    store.create_task(next_task)
    second=service.generate(next_task,task.analysis_ids,title='新版简报',base_report_id=result.report_ids[0])
    assert not second.validation_summary['layout']['include_quality_and_limitations']
    assert not second.validation_summary['layout']['include_sources']


def test_old_four_switch_snapshot_inherits_new_defaults_without_restoring_removed_chapters(reports,monkeypatch):
    store,service,_,source,task=reports
    register=store.register_artifact
    def old_register(path,*args,**kwargs):
        if path.name=='manifest.json':
            value=json.loads(path.read_text(encoding='utf-8'))
            value['presentation']['layout'].pop('include_quality_and_limitations')
            value['presentation']['layout'].pop('include_sources')
            path.write_text(json.dumps(value,ensure_ascii=False),encoding='utf-8')
        return register(path,*args,**kwargs)
    monkeypatch.setattr(store,'register_artifact',old_register)
    first=service.generate(task,task.analysis_ids,base_report_id=source.report_ids[0],layout={'include_metrics_summary':False,'include_findings':False})
    deliver(store,task,first)
    monkeypatch.setattr(store,'register_artifact',register)
    next_task=TaskSpec(session_id=task.session_id,request='删除质量模块')
    store.create_task(next_task)
    base=service.read_report(next_task,first.report_ids[0])
    assert base['presentation']['layout']['include_quality_and_limitations']
    assert not base['presentation']['layout']['include_metrics_summary']
    second=service.generate(next_task,task.analysis_ids,base_report_id=first.report_ids[0],layout={'include_quality_and_limitations':False})
    assert not second.validation_summary['layout']['include_metrics_summary']
    assert not second.validation_summary['layout']['include_findings']
    assert not second.validation_summary['layout']['include_quality_and_limitations']


def test_noop_revision_does_not_publish_timestamp_only_new_file(reports):
    store,service,_,source,task=reports
    with pytest.raises(StoreError) as error:
        service.generate(task,task.analysis_ids,base_report_id=source.report_ids[0],
                         layout={'include_metrics_summary':True})
    assert error.value.code=='report_unchanged'
    assert not [a for a in store.list_artifacts(task.task_id) if a.kind.startswith('report')]
