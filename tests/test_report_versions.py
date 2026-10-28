"""Report history remains available across revisions, reloads, and failures."""
import json

from fastapi.testclient import TestClient
from bs4 import BeautifulSoup
from openpyxl import load_workbook
import pytest

from src.api import create_app
from src.config import Settings
from src.data import DataService
from src.reporting import ReportService
from src.schemas import TaskSpec
from src.storage import Store, file_hash


@pytest.fixture
def history(tmp_path):
    store = Store(Settings(workspace=tmp_path, _env_file=None))
    sid = store.new_session()['session_id']
    path = tmp_path / 'KB' / 'data.csv'
    path.write_text('value\n1\n3\n', encoding='utf-8')
    fid = store.refresh_kb()[0]['file_id']
    task = TaskSpec(session_id=sid, request='原版', file_ids=[fid])
    store.create_task(task)
    analysis = DataService(store).analyze(task, fid, 'aggregate', aggregations={'value': 'mean'})
    service = ReportService(store)
    original = service.generate(task, [analysis.analysis_id], title='原版', layout={'include_metrics_summary': True})
    store.update_task(task.task_id, status='succeeded', result=original.model_dump())
    revision = TaskSpec(session_id=sid, request='修改标题', analysis_ids=[analysis.analysis_id])
    store.create_task(revision)
    return store, service, task, original, revision, analysis


def assert_no_metrics(store, result):
    for rid in result.report_ids:
        path = store.resolve_artifact(rid)
        if path.suffix == '.html':
            soup = BeautifulSoup(path.read_text(encoding='utf-8'), 'html.parser')
            assert not soup.find('h2', string='指标汇总')
            assert not soup.select('.metric-value')
        else:
            book = load_workbook(path, read_only=True)
            try:
                assert '指标汇总' not in book.sheetnames
                assert '结果表_1_1' not in book.sheetnames
            finally:
                book.close()
    manifest = json.loads(store.resolve_artifact(result.validation_summary['manifest_id']).read_text(encoding='utf-8'))
    assert manifest['metric_bindings']  # Data and provenance remain available.
    assert manifest['validation']['metrics_summary_included'] is False


def test_new_report_and_legacy_revision_default_to_no_metrics_summary(history):
    store, service, task, original, revision, analysis = history
    hashes = {rid: file_hash(store.resolve_artifact(rid)) for rid in original.report_ids}
    fresh = service.generate(revision, [analysis.analysis_id])
    assert_no_metrics(store, fresh)
    edited = service.generate(revision, [analysis.analysis_id], title='修改版', base_report_id=original.report_ids[0])
    assert_no_metrics(store, edited)
    assert edited.validation_summary['layout_changes'] == {'include_metrics_summary': False}
    assert hashes == {rid: file_hash(store.resolve_artifact(rid)) for rid in original.report_ids}


def test_session_history_and_downloads_keep_all_versions_after_restart(history):
    store, service, task, original, revision, analysis = history
    originals = {rid: store.resolve_artifact(rid).read_bytes() for rid in original.report_ids}
    edited = service.generate(revision, [analysis.analysis_id], title='修改版', base_report_id=original.report_ids[0])
    store.update_task(revision.task_id, status='succeeded', result=edited.model_dump())
    # Recreate the API to ensure catalog reconstruction uses persistent records.
    with TestClient(create_app(store.settings)) as client:
        versions = client.get(f'/api/v1/sessions/{task.session_id}').json()['report_versions']
        assert [v['title'] for v in versions] == ['修改版', '原版']
        assert [v['number'] for v in versions] == [2, 1]
        assert versions[0]['base_number'] == 1
        assert versions[0]['revision_of']['report_id'] == original.report_ids[0]
        assert versions[1]['revision_of'] is None
        assert {a['artifact_id'] for v in versions for a in v['artifacts']} == set(original.report_ids + edited.report_ids)
        assert len({store.resolve_artifact(rid).parent for rid in original.report_ids + edited.report_ids}) == 2
        for rid, data in originals.items():
            assert client.get(f'/api/v1/artifacts/{rid}/download').content == data
        for rid in edited.report_ids:
            assert client.get(f'/api/v1/artifacts/{rid}/download').status_code == 200
        other_sid = client.post('/api/v1/sessions').json()['session_id']
        assert client.get(f'/api/v1/sessions/{other_sid}').json()['report_versions'] == []


def test_failed_revision_cleanup_keeps_original_files_and_catalog(history, monkeypatch):
    store, service, task, original, revision, analysis = history
    hashes = {rid: file_hash(store.resolve_artifact(rid)) for rid in original.report_ids}
    def broken(*args):
        raise RuntimeError('simulated validation failure')
    monkeypatch.setattr(service, '_validate', broken)
    with pytest.raises(RuntimeError, match='simulated'):
        service.generate(revision, [analysis.analysis_id], title='失败版本', base_report_id=original.report_ids[0])
    store.update_task(revision.task_id, status='failed')
    assert hashes == {rid: file_hash(store.resolve_artifact(rid)) for rid in original.report_ids}
    assert not store.list_artifacts(revision.task_id)
    with TestClient(create_app(store.settings)) as client:
        versions = client.get(f'/api/v1/sessions/{task.session_id}').json()['report_versions']
        assert len(versions) == 1 and versions[0]['title'] == '原版'
