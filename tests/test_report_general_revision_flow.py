"""Independent acceptance probes for general, source-bound report patches.

These intentionally use non-Top-N plots and new verified analyses. They check
the invariant the user asked for: edit the requested part of the chosen base,
preserve everything else, and keep every previous version.
"""
from __future__ import annotations

import json

import pytest
from bs4 import BeautifulSoup
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl import load_workbook

from src.chart_export import configure_chart_fonts, export_figure
from src.config import Settings
from src.data import DataService
from src.reporting import ReportService
from src.schemas import TaskSpec
from src.storage import Store, StoreError, file_hash


def manifest(store, result):
    return json.loads(store.resolve_artifact(result.validation_summary['manifest_id']).read_text(encoding='utf-8'))


def html_soup(store, result):
    report_ids = result['report_ids'] if isinstance(result, dict) else result.report_ids
    path = next(store.resolve_artifact(rid) for rid in report_ids
                if store.get_artifact(rid).name.endswith('.html'))
    return BeautifulSoup(path.read_text(encoding='utf-8'), 'html.parser')


def register_plot(store, task, analysis, name, kind='line'):
    configure_chart_fonts()
    figure = Figure(figsize=(5, 3)); FigureCanvasAgg(figure)
    axes = figure.subplots()
    if kind == 'scatter':
        axes.scatter([1, 2, 3, 4], [10, 20, 30, 40], label='Observed')
    elif kind == 'multi':
        axes.plot([1, 2], [30, 40], label='Delivered amount')
        axes.plot([1, 2], [15, 20], label='Reference amount')
        axes.legend()
    else:
        axes.plot([1, 2, 3, 4], [2, 4, 3, 5], label='Unchanged series')
    axes.set_title(name)
    path = store.safe_path(f'script_output/{task.session_id}/{task.task_id}/{name}.svg')
    path.parent.mkdir(parents=True, exist_ok=True)
    export_figure(figure, path, Figure.savefig)
    png = path.with_suffix('.png'); figure.savefig(png)
    artifact = store.register_artifact(path, task.task_id, task.session_id, 'chart', list(analysis.input_versions))
    store.register_artifact(png, task.task_id, task.session_id, 'chart_preview', [artifact.artifact_id])
    analysis.chart_ids.append(artifact.artifact_id)
    store.save_analysis(analysis)
    return artifact.artifact_id


@pytest.fixture
def general_base(tmp_path):
    store = Store(Settings(workspace=tmp_path, _env_file=None))
    sid = store.new_session()['session_id']
    (tmp_path/'KB'/'events.csv').write_text(
        'category,amount,status\nA,10,pending\nB,20,pending\nC,30,delivered\nD,40,delivered\n', encoding='utf-8')
    fid = store.refresh_kb()[0]['file_id']
    base_task = TaskSpec(session_id=sid, request='Create overview report', file_ids=[fid])
    store.create_task(base_task)
    data = DataService(store)
    analysis = data.analyze(base_task, fid, 'aggregate', aggregations={'amount': 'mean'})
    full = data.analyze(base_task, fid, 'rows')
    scatter = register_plot(store, base_task, full, 'Original scatter', 'scatter')
    line = register_plot(store, base_task, full, 'Preserved line')
    service = ReportService(store)
    result = service.generate(base_task, [analysis.analysis_id, full.analysis_id], title='Business report',
        narrative={'summary': 'Average amount is {{metric:' + analysis.analysis_id + ':amount_mean}}.',
                   'sections': [{'heading': 'Scope', 'text': 'All authorized events.',
                                 'metric_refs': [{'analysis_id': analysis.analysis_id, 'metric_key': 'amount_mean'}],
                                 'chart_ids': [scatter], 'table_ids': [full.table_ids[0]]},
                                {'heading': 'Recommendations', 'text': 'Keep this exact paragraph.'}]},
        formats=['html', 'xlsx'])
    store.update_task(base_task.task_id, status='succeeded', result=result.model_dump())
    hashes = {item.artifact_id: file_hash(store.resolve_artifact(item.artifact_id))
              for item in store.list_artifacts(base_task.task_id)}
    return store, service, data, fid, analysis, full, scatter, line, result, hashes


def next_task(store, fid, request='Modify the previous report'):
    task = TaskSpec(session_id=store.get_file(fid, '')['session_id'] or store.list_sessions()[0]['session_id'],
                    request=request, file_ids=[fid])
    store.create_task(task)
    return task


def test_new_filtered_analysis_replaces_only_one_nonbar_chart_and_requested_conclusion(general_base):
    store, service, data, fid, average, full, scatter, line, base, hashes = general_base
    task = next_task(store, fid)
    supplement = data.analyze(task, fid, 'aggregate', aggregations={'amount': 'mean'},
                              filters=[{'column': 'status', 'op': 'eq', 'value': 'delivered'}])
    replacement = register_plot(store, task, supplement, 'Delivered comparison', 'multi')
    extra = register_plot(store, task, supplement, 'Unrequested supplemental plot')
    patch = {'summary': 'Delivered average is {{metric:' + supplement.analysis_id + ':amount_mean}}.',
             'components': [{'kind': 'chart', 'operation': 'replace', 'target_id': scatter,
                             'artifact_id': replacement, 'analysis_id': supplement.analysis_id}]}
    revised = service.generate(task, [average.analysis_id, full.analysis_id, supplement.analysis_id],
                               base_report_id=base.report_ids[0], revision_patch=patch)
    before, after = manifest(store, base), manifest(store, revised)
    assert after['presentation']['summary'] == 'Delivered average is 35.0.'
    assert after['presentation']['sections'] == before['presentation']['sections']
    assert after['presentation']['findings'] == before['presentation']['findings']
    old_html, new_html = html_soup(store, base), html_soup(store, revised)
    assert str(new_html.select_one(f'figure[data-chart-card="{line}"]')) == str(old_html.select_one(f'figure[data-chart-card="{line}"]'))
    assert len(new_html.select('figure.chart-card')) == 2
    assert new_html.select_one(f'figure[data-chart-card="{replacement}"]')
    assert not new_html.select_one(f'figure[data-chart-card="{extra}"]')
    assert not new_html.select_one(f'figure[data-chart-card="{scatter}"]')
    assert not new_html.find('h2', string='结果明细')
    assert after['analyses'][:len(before['analyses'])] == before['analyses']
    assert len(after['analyses']) == len(before['analyses']) + 1
    assert hashes == {aid: file_hash(store.resolve_artifact(aid)) for aid in hashes}
    store.update_task(task.task_id, status='succeeded', result=revised.model_dump())
    restored = service.read_report(task, revised.report_ids[0])
    component = next(item for item in restored['component_state']['charts'] if item['component_id'] == scatter)
    assert component['artifact_id'] == replacement and component['visible']


def test_add_one_requested_table_does_not_reveal_existing_hidden_result_tables(general_base):
    store, service, data, fid, average, full, _, _, base, _ = general_base
    task = next_task(store, fid)
    rows = data.analyze(task, fid, 'rows', filters=[{'column': 'status', 'op': 'eq', 'value': 'delivered'}])
    added = rows.table_ids[0]
    changed = service.generate(task, [average.analysis_id, full.analysis_id, rows.analysis_id],
        base_report_id=base.report_ids[0], revision_patch={'components': [
            {'kind': 'table', 'operation': 'add', 'artifact_id': added,
             'analysis_id': rows.analysis_id, 'visible': True}]})
    soup = html_soup(store, changed)
    assert added in soup.get_text()
    assert all(tid not in soup.get_text() for tid in average.table_ids + full.table_ids)
    xlsx = next(store.resolve_artifact(rid) for rid in changed.report_ids if store.get_artifact(rid).name.endswith('.xlsx'))
    book = load_workbook(xlsx, read_only=True)
    try:
        names = [name for name in book.sheetnames if name.startswith('结果表_')]
        assert len(names) == 1 and book[names[0]].max_row == 3
    finally:
        book.close()
    store.update_task(task.task_id, status='succeeded', result=changed.model_dump())
    followup = next_task(store, fid, 'Only rename the report')
    renamed = service.generate(followup, [average.analysis_id, full.analysis_id, rows.analysis_id],
                               title='Renamed business report', base_report_id=changed.report_ids[0])
    assert added in html_soup(store, renamed).get_text()
    assert all(tid not in html_soup(store, renamed).get_text() for tid in average.table_ids + full.table_ids)


def test_hide_conclusion_then_insert_move_and_remove_sections_inherits_current_version(general_base):
    store, service, _, fid, average, full, _, _, base, _ = general_base
    task = next_task(store, fid)
    hidden = service.generate(task, [average.analysis_id, full.analysis_id], base_report_id=base.report_ids[0],
        layout={'include_conclusion': False}, revision_patch={'sections': [
            {'operation': 'add', 'position': 1, 'section': {'heading': 'Actions', 'text': 'Act on the observed pattern.'}},
            {'operation': 'move', 'target_heading': 'Recommendations', 'position': 0}]})
    assert not html_soup(store, hidden).select_one('[data-section="conclusion"]')
    assert [section['heading'] for section in manifest(store, hidden)['presentation']['sections']] == ['Recommendations', 'Scope', 'Actions']
    store.update_task(task.task_id, status='succeeded', result=hidden.model_dump())
    followup = next_task(store, fid)
    changed = service.generate(followup, [average.analysis_id, full.analysis_id], base_report_id=hidden.report_ids[0],
        revision_patch={'sections': [{'operation': 'remove', 'target_heading': 'Scope'}]})
    assert not html_soup(store, changed).select_one('[data-section="conclusion"]')
    assert [section['heading'] for section in manifest(store, changed)['presentation']['sections']] == ['Recommendations', 'Actions']
    xlsx = next(store.resolve_artifact(rid) for rid in changed.report_ids if store.get_artifact(rid).name.endswith('.xlsx'))
    book = load_workbook(xlsx, read_only=True)
    try:
        assert '结论' not in dict(book['说明与口径'].values)
    finally:
        book.close()


def test_patch_branched_from_explicit_older_base_does_not_import_latest_changes(general_base):
    store, service, _, fid, average, full, _, _, base, hashes = general_base
    first = next_task(store, fid)
    one = service.generate(first, [average.analysis_id, full.analysis_id], title='Latest modified title', base_report_id=base.report_ids[0])
    store.update_task(first.task_id, status='succeeded', result=one.model_dump())
    branch = next_task(store, fid, 'Use original version, only change conclusion')
    two = service.generate(branch, [average.analysis_id, full.analysis_id], base_report_id=base.report_ids[0],
        revision_patch={'summary': 'A shorter explanation of the same evidence.'})
    assert manifest(store, two)['presentation']['title'] == 'Business report'
    assert manifest(store, two)['revision_of']['report_id'] == base.report_ids[0]
    assert hashes == {aid: file_hash(store.resolve_artifact(aid)) for aid in hashes}


def test_text_only_chapter_patch_preserves_its_metric_and_artifact_references(general_base):
    store, service, _, fid, average, full, _, _, base, _ = general_base
    task = next_task(store, fid)
    before = manifest(store, base)['presentation']['sections'][0]
    revised = service.generate(task, [average.analysis_id, full.analysis_id], base_report_id=base.report_ids[0],
        revision_patch={'sections': [{'operation': 'update', 'target_heading': 'Scope',
                                    'section': {'text': 'A shorter scope explanation.'}}]})
    after = manifest(store, revised)['presentation']['sections'][0]
    assert after == {**before, 'text': 'A shorter scope explanation.'}


@pytest.mark.parametrize('issue', ['unknown_target', 'wrong_owner', 'unreferenced_extra', 'numeric_claim'])
def test_invalid_generic_patches_do_not_publish(general_base, issue):
    store, service, data, fid, average, full, scatter, _, base, _ = general_base
    task = next_task(store, fid)
    supplement = data.analyze(task, fid, 'rows')
    replacement = register_plot(store, task, supplement, 'Replacement scatter', 'scatter')
    ids = [average.analysis_id, full.analysis_id, supplement.analysis_id]
    patch = {'components': [{'kind': 'chart', 'operation': 'replace', 'target_id': scatter,
                             'artifact_id': replacement, 'analysis_id': supplement.analysis_id}]}
    if issue == 'unknown_target':
        patch['components'][0]['target_id'] = 'not-an-existing-component'
    elif issue == 'wrong_owner':
        patch['components'][0]['analysis_id'] = average.analysis_id
    elif issue == 'unreferenced_extra':
        patch = {'summary': 'Only a wording change.'}
    else:
        patch['summary'] = 'The delivered average is 9999.'
    try:
        result = service.generate(task, ids, base_report_id=base.report_ids[0], revision_patch=patch)
    except StoreError:
        pass
    else:
        assert result.status != 'completed', result
    assert not any(item.kind == 'report' for item in store.list_artifacts(task.task_id))


@pytest.mark.parametrize('operation', ['add', 'replace'])
def test_duplicate_active_chart_artifacts_cannot_duplicate_svg_dom_ids(general_base, operation):
    store, service, _, fid, average, full, scatter, line, base, _ = general_base
    task = next_task(store, fid)
    edit = {'kind': 'chart', 'operation': operation, 'artifact_id': line, 'analysis_id': full.analysis_id}
    if operation == 'replace':
        edit['target_id'] = scatter
    try:
        result = service.generate(task, [average.analysis_id, full.analysis_id],
            base_report_id=base.report_ids[0], revision_patch={'components': [edit]})
    except StoreError:
        pass
    else:
        if result.status == 'completed':
            soup = html_soup(store, result)
            ids = [item['id'] for item in soup.select('[id]')]
            assert len(ids) == len(set(ids)), 'Repeated artifacts generated duplicate active SVG IDs'


def test_supplemental_material_evidence_does_not_implicitly_append_its_raw_excerpt(general_base):
    from unittest.mock import Mock
    from src.agents import TaskRuntime, analyst_tools

    store, service, _, fid, average, full, _, _, base, _ = general_base
    path = store.root/'KB'/'context.html'
    path.write_text('<html><body>RAW_EXTRA_ONLY_NOT_REQUESTED_IN_REPORT</body></html>', encoding='utf-8')
    file = store.register_file(path, path.name, None, source='kb')
    task = next_task(store, fid, 'Only add a document-length sentence to the conclusion')
    task.file_ids.append(file['file_id'])
    store.update_task(task.task_id, spec=task.model_dump())
    runtime = TaskRuntime(task, store, Mock(), lambda: False)
    tools = {tool.name: tool for tool in analyst_tools(runtime)}
    prepared = tools['prepare_material'].invoke({'file_id': file['file_id']})
    assert prepared['analysis_id']
    revised = service.generate(task, [average.analysis_id, full.analysis_id, prepared['analysis_id']],
        base_report_id=base.report_ids[0], revision_patch={
            'summary': 'Document characters: {{metric:' + prepared['analysis_id'] + ':source_characters}}.'})
    assert revised.status == 'completed'
    soup = html_soup(store, revised)
    assert 'RAW_EXTRA_ONLY_NOT_REQUESTED_IN_REPORT' not in soup.get_text()
    assert not soup.find('h2', string='材料原文摘录')


def test_pinned_chapter_scope_rejects_changes_to_an_unrelated_conclusion(general_base):
    from unittest.mock import Mock
    from src.agents import TaskRuntime, report_tools

    store, _, _, fid, _, _, _, _, base, _ = general_base
    task = next_task(store, fid, 'Only shorten Scope')
    runtime = TaskRuntime(task, store, Mock(), lambda: False)
    runtime.prepare_revision(base.report_ids[0], targets=['section:Scope'])
    tools = {tool.name: tool for tool in report_tools(runtime)}
    result = tools['revise_report'].invoke({'revision_patch': {'summary': 'Changed outside the pinned chapter.'}})
    assert result['status'] != 'completed', result
    assert not any(item.kind == 'report' for item in store.list_artifacts(task.task_id))


def test_formal_plan_cannot_silently_switch_to_another_base_before_supplement(general_base):
    from unittest.mock import Mock
    from src.agents import TaskRuntime

    store, service, _, fid, average, full, _, _, base, _ = general_base
    first = next_task(store, fid)
    newer = service.generate(first, [average.analysis_id, full.analysis_id], title='Newer title', base_report_id=base.report_ids[0])
    store.update_task(first.task_id, status='succeeded', result=newer.model_dump())
    task = next_task(store, fid)
    runtime = TaskRuntime(task, store, Mock(), lambda: False)
    runtime.prepare_revision(base.report_ids[0], targets=['summary'])
    with pytest.raises(StoreError):
        runtime.prepare_revision(newer.report_ids[0])


def test_formal_visibility_plan_can_hide_conclusion_without_rewriting_other_fields(general_base):
    from unittest.mock import Mock
    from src.agents import TaskRuntime, report_tools

    store, _, _, fid, _, _, _, _, base, _ = general_base
    task = next_task(store, fid, 'Remove only the conclusion')
    runtime = TaskRuntime(task, store, Mock(), lambda: False)
    runtime.prepare_revision(base.report_ids[0], targets=['layout:include_conclusion'])
    tools = {tool.name: tool for tool in report_tools(runtime)}
    result = tools['revise_report'].invoke({'layout_changes': {'include_conclusion': False}})
    assert result['status'] == 'completed', result
    assert not html_soup(store, result).select_one('[data-section="conclusion"]')


def test_plan_for_one_chapter_does_not_grant_moves_of_a_different_chapter(general_base):
    from unittest.mock import Mock
    from src.agents import TaskRuntime, report_tools

    store, _, _, fid, _, _, _, _, base, _ = general_base
    task = next_task(store, fid, 'Only edit Scope')
    runtime = TaskRuntime(task, store, Mock(), lambda: False)
    runtime.prepare_revision(base.report_ids[0], targets=['section:Scope'])
    tools = {tool.name: tool for tool in report_tools(runtime)}
    result = tools['revise_report'].invoke({'revision_patch': {'sections': [
        {'operation': 'move', 'target_heading': 'Recommendations', 'position': 0}]}})
    assert result['status'] != 'completed', result
    assert not any(item.kind == 'report' for item in store.list_artifacts(task.task_id))
