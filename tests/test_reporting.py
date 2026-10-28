import json
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from src.config import Settings
from src.data import DataService
from src.reporting import ReportService
from src.schemas import AnalysisResult, TaskSpec
from src.storage import Store, StoreError


@pytest.fixture
def context(tmp_path):
    store = Store(Settings(workspace=tmp_path, _env_file=None))
    sid = store.new_session()['session_id']
    path = store.root / 'KB' / 'sample.csv'
    path.write_text('id,value,other\n001,2,5\nNA,4,7\n003,9,11\n', encoding='utf-8')
    fid = store.refresh_kb()[0]['file_id']
    task = TaskSpec(session_id=sid, request='Analyze and report', file_ids=[fid])
    store.create_task(task)
    data = DataService(store)
    mean = data.analyze(task, fid, 'aggregate', aggregations={'value': ['mean', 'sum']})
    return store, task, fid, mean


def outputs(store, result):
    return {store.get_artifact(aid).name: store.resolve_artifact(aid) for aid in result.report_ids}


def test_both_formats_bind_real_metrics_and_sources(context):
    store, task, fid, analysis = context
    second = DataService(store).analyze(task, fid, 'correlation', columns=['value', 'other'], chart='bar')
    narrative = {'summary': '均值为{{metric:' + analysis.analysis_id + ':value_mean}}。',
        'sections': [{'heading': '核心结果', 'text': '结果来自所选完整数据。',
            'metric_refs': [{'analysis_id': analysis.analysis_id, 'metric_key': 'value_sum', 'label': '合计'}],
            'table_ids': analysis.table_ids, 'chart_ids': second.chart_ids}]}
    result = ReportService(store).generate(task, [analysis.analysis_id, second.analysis_id], narrative=narrative,
                                          layout={'include_metrics_summary': True, 'include_tables': True})
    assert result.status == 'completed'
    files = outputs(store, result)
    assert set(files) == {'report.html', 'report.xlsx'}
    html = files['report.html'].read_text(encoding='utf-8')
    assert 'data-ds-chart=' in html and 'data:image/png;base64,' not in html
    assert len(BeautifulSoup(html, 'html.parser').select('svg[data-ds-chart]')) == 1
    assert '均值为5.0。' in html
    assert store.get_file(fid, task.session_id)['hash'] in html
    assert 'https://' not in html
    book = load_workbook(files['report.xlsx'], data_only=False)
    try:
        metrics = {(r[0], r[1]): r[3] for r in list(book['指标汇总'].values)[1:]}
        assert metrics[(analysis.analysis_id, 'value_mean')] == '5.0'
        assert metrics[(analysis.analysis_id, 'value_sum')] == '15'
        assert book['结果表_1_1'].max_row == 2
        assert len(book['图表']._images) == 1
    finally:
        book.close()
    manifest = json.loads(store.resolve_artifact(result.validation_summary['manifest_id']).read_text(encoding='utf-8'))
    assert manifest['timezone'] == 'Asia/Shanghai'
    assert manifest['sources'][0]['file_id'] == fid
    assert manifest['analyses'][0]['metric_definitions']
    assert manifest['models']['main'] == store.settings.main_model
    assert 'api_key' not in json.dumps(manifest)


def test_text_ids_formula_injection_xml_cleanup_and_full_tables(context):
    store, task, fid, _ = context
    source = store.resolve_file(fid, task.session_id)
    # Re-register a new immutable source version rather than mutating an analysis.
    source.write_text('id,text,value\n001,=1+1,-2\nNA,+SUM(A1),-3\n003,@evil,4\n004,-formula,5\n005,</td><script>alert</script>,6\n006,bad\x01text,7\n' +
        ''.join(f'{100+i},row,{i}\n' for i in range(250)), encoding='utf-8')
    new_file = store.refresh_kb()[0]
    task.file_ids = [new_file['file_id']]
    store.update_task(task.task_id, spec=task.model_dump())
    analysis = DataService(store).analyze(task, new_file['file_id'], 'rows')
    result = ReportService(store).generate(task, [analysis.analysis_id], title='<img src=x onerror=alert()>', layout={'include_tables': True})
    files = outputs(store, result)
    html = files['report.html'].read_text(encoding='utf-8')
    soup = BeautifulSoup(html, 'html.parser')
    assert not soup.find('script') and not soup.find('img')
    assert '&lt;script&gt;' in html and '&lt;img' in html
    assert '此处为预览' in html and '256' in html
    assert 'badtext' in html
    assert result.validation_summary['removed_xml_characters'] > 0
    book = load_workbook(files['report.xlsx'], data_only=False)
    try:
        sheet = book['结果表_1_1']
        assert sheet.max_row == 257
        assert sheet['A2'].value == '001' and sheet['A3'].value == 'NA'
        assert sheet['B2'].value == '=1+1' and sheet['B2'].data_type == 's'
        assert sheet['B3'].value == '+SUM(A1)' and sheet['B4'].value == '@evil'
        assert sheet['B5'].value == '-formula'
        assert sheet['C2'].value == -2 and sheet['C2'].data_type == 'n'
        assert all(c.data_type != 'f' for row in sheet for c in row)
    finally:
        book.close()


@pytest.mark.parametrize('narrative', [
    {'summary': '平均是999。'},
    {'summary': '平均{{metric:missing:mean}}'},
    {'sections': [{'metric_refs': [{'analysis_id': 'missing', 'metric_key': 'mean'}]}]},
])
def test_unknown_or_unbound_numbers_request_analysis_without_outputs(context, narrative):
    store, task, _, analysis = context
    result = ReportService(store).generate(task, [analysis.analysis_id], narrative=narrative)
    assert result.status == 'need_analysis' and result.supplement_request
    assert not result.report_ids
    assert not [a for a in store.list_artifacts(task.task_id) if a.kind.startswith('report')]
    assert not list((store.root / 'report').rglob('*.html'))


@pytest.mark.parametrize('issue', ['other_session', 'stale_source', 'tampered_artifact', 'unauthorized', 'wrong_kind', 'wrong_chart', 'out_of_scope'])
def test_invalid_sources_and_artifacts_rejected(context, issue):
    store, task, fid, analysis = context
    if issue == 'other_session':
        sid = store.new_session()['session_id']
        task = TaskSpec(session_id=sid, request='Other', file_ids=[fid], analysis_ids=[analysis.analysis_id])
        store.create_task(task)
    elif issue == 'stale_source':
        store.resolve_file(fid, task.session_id).write_text('value\n999\n')
    elif issue == 'tampered_artifact':
        store.resolve_artifact(analysis.table_ids[0]).write_text('bad')
    elif issue == 'unauthorized':
        task = TaskSpec(session_id=task.session_id, request='New task', file_ids=[fid])
        store.create_task(task)
    elif issue == 'wrong_kind':
        analysis.chart_ids = analysis.table_ids
        analysis.table_ids = []
        store.save_analysis(analysis)
    elif issue == 'wrong_chart':
        path = store.safe_path(f'script_output/{task.session_id}/{task.task_id}/bad.png')
        path.write_bytes(b'not-a-png')
        artifact = store.register_artifact(path, task.task_id, task.session_id, 'chart')
        analysis.chart_ids = [artifact.artifact_id]
        store.save_analysis(analysis)
    else:
        (store.root / 'KB' / 'different.csv').write_text('x\n1\n')
        task.file_ids = [x['file_id'] for x in store.refresh_kb() if x['name'] == 'different.csv']
    with pytest.raises(StoreError):
        ReportService(store).generate(task, [analysis.analysis_id])
    assert not list((store.root / 'report').rglob('report.*'))


def test_versions_formats_and_metrics_only(context):
    store, task, _, analysis = context
    # A valid analysis need not have any result table or chart.
    analysis.table_ids = []
    analysis.metrics['precise'] = 0.12345678901234568
    store.save_analysis(analysis)
    service = ReportService(store)
    first = service.generate(task, [analysis.analysis_id], formats=['html'])
    original = store.resolve_artifact(first.report_ids[0]).read_bytes()
    second = service.generate(task, [analysis.analysis_id], formats=['xlsx'], layout={'include_metrics_summary': True})
    assert first.report_ids != second.report_ids
    assert store.resolve_artifact(first.report_ids[0]).read_bytes() == original
    assert store.resolve_artifact(first.report_ids[0]).parent != store.resolve_artifact(second.report_ids[0]).parent
    book = load_workbook(store.resolve_artifact(second.report_ids[0]))
    try:
        row = next(r for r in book['指标汇总'].values if r[1] == 'precise')
        assert row[3] == '0.12345678901234568'
        assert row[4] == '0.12345678901234568'
    finally:
        book.close()
    assert ReportService(store).generate(task, []).status == 'need_analysis'
    for formats in [[], ['pdf']]:
        with pytest.raises(StoreError, match='formats'):
            service.generate(task, [analysis.analysis_id], formats=formats)


def test_render_failure_does_not_publish_or_register_partial_reports(context, monkeypatch):
    store, task, _, analysis = context
    service = ReportService(store)
    monkeypatch.setattr(service, '_validate', lambda *a: (_ for _ in ()).throw(RuntimeError('forced validation failure')))
    with pytest.raises(RuntimeError, match='forced'):
        service.generate(task, [analysis.analysis_id])
    assert not list((store.root / 'report').rglob('report.*'))
    assert not list((store.root / 'script_output').rglob('.report_staging_*'))
    assert not [a for a in store.list_artifacts(task.task_id) if a.kind.startswith('report')]


def test_registration_failure_rolls_back_reports(context, monkeypatch):
    store, task, _, analysis = context
    original = store.register_artifact
    def fail_manifest(path, *args, **kwargs):
        if Path(path).suffix == '.json':
            raise RuntimeError('forced registration failure')
        return original(path, *args, **kwargs)
    monkeypatch.setattr(store, 'register_artifact', fail_manifest)
    with pytest.raises(RuntimeError, match='registration'):
        ReportService(store).generate(task, [analysis.analysis_id])
    assert not list((store.root / 'report').rglob('report.*'))
    assert not [a for a in store.list_artifacts(task.task_id) if a.kind.startswith('report')]


def test_reuse_only_explicit_or_successfully_registered_same_session_analysis(context):
    store, task, fid, analysis = context
    next_task = TaskSpec(session_id=task.session_id, request='Reuse results', file_ids=[fid], analysis_ids=[analysis.analysis_id])
    store.create_task(next_task)
    assert ReportService(store).generate(next_task, [analysis.analysis_id], formats=['html']).status == 'completed'
    next_task.analysis_ids = []
    store.update_task(task.task_id, status='succeeded', result={'analysis_ids': [analysis.analysis_id]})
    assert ReportService(store).generate(next_task, [analysis.analysis_id], formats=['xlsx']).status == 'completed'
    store.update_task(task.task_id, result={})
    waiting_parent = TaskSpec(session_id=task.session_id, request='Clarify audience', file_ids=[fid])
    store.create_task(waiting_parent)
    store.update_task(waiting_parent.task_id, status='waiting_user', result={'analysis_ids': [analysis.analysis_id]})
    next_task.parent_task_id = waiting_parent.task_id
    with pytest.raises(StoreError, match='authorized'):
        ReportService(store).generate(next_task, [analysis.analysis_id])


def test_material_excerpt_is_explicitly_nonstatistical_and_source_bound(context):
    store, task, fid, _ = context
    material = AnalysisResult(task_id=task.task_id, session_id=task.session_id,
        input_versions={fid: store.get_file(fid, task.session_id)['hash']},
        metrics={}, kb_refs=[dict(source_type='material', file_id=fid, location='extracted text',
            text='<b>实际材料摘录</b>', truncated=True)])
    store.save_analysis(material)
    result = ReportService(store).generate(task, [material.analysis_id])
    files = outputs(store, result)
    html = files['report.html'].read_text(encoding='utf-8')
    assert not BeautifulSoup(html, 'html.parser').find('h2', string='指标汇总')
    assert '未进行业务统计' in html
    assert '原文已截断' in html and '&lt;b&gt;实际材料摘录&lt;/b&gt;' in html
    book = load_workbook(files['report.xlsx'])
    try:
        assert book['材料摘录']['D2'].value == '<b>实际材料摘录</b>'
        assert '已截断' in book['材料摘录']['E2'].value
    finally:
        book.close()


@pytest.mark.parametrize('changed', ['source', 'analysis'])
def test_changes_during_rendering_are_rejected_before_publish(context, monkeypatch, changed):
    store, task, fid, analysis = context
    service = ReportService(store)
    validate = service._validate
    source = store.resolve_file(fid, task.session_id)
    def change_after_validate(*args):
        validate(*args)
        if changed == 'source':
            source.write_text('value\n999\n')
        else:
            analysis.metrics['value_mean'] = 999
            store.save_analysis(analysis)
    monkeypatch.setattr(service, '_validate', change_after_validate)
    with pytest.raises(StoreError):
        service.generate(task, [analysis.analysis_id])
    assert not list((store.root / 'report').rglob('report.*'))
    assert not [a for a in store.list_artifacts(task.task_id) if a.kind.startswith('report')]


@pytest.mark.parametrize('text', ['均值为九百九十九。', '均值为九。', '约三', '占比百分之九十', '均值为９９９', '九点五', '负九'])
def test_numeric_claims_including_chinese_require_bindings(context, text):
    store, task, _, analysis = context
    result = ReportService(store).generate(task, [analysis.analysis_id], narrative={'summary': text})
    assert result.status == 'need_analysis'
    assert not result.report_ids


def test_title_requires_binding_and_qualitative_chinese_is_allowed(context):
    store, task, _, analysis = context
    service = ReportService(store)
    assert service.generate(task, [analysis.analysis_id], title='均值为999').status == 'need_analysis'
    result = service.generate(task, [analysis.analysis_id], title='均值{{metric:' + analysis.analysis_id + ':value_mean}}',
        narrative={'summary': '统计口径一致，统一展示；一般按已验证结果整理，一方面保留来源，另一方面说明局限。观点负面仅为定性意见。'}, formats=['html'])
    assert result.status == 'completed'


@pytest.mark.parametrize('narrative,reason_code', [
    ({'summary': '平均为999'}, 'unbound_prose'),
    ({'summary': '均值{{metric:malformed}}'}, 'invalid_placeholder'),
    ({'summary': '均值{{metric:missing:mean}}'}, 'missing_metric'),
    ({'sections': [{'metric_refs': [{'analysis_id': 'missing', 'metric_key': 'mean'}]}]}, 'missing_metric'),
])
def test_report_validation_reason_distinguishes_revision_from_missing_data(context, narrative, reason_code):
    store, task, _, analysis = context
    result = ReportService(store).generate(task, [analysis.analysis_id], narrative=narrative)
    assert result.status == 'need_analysis'
    assert result.validation_summary['reason_code'] == reason_code
    assert result.analysis_ids == [analysis.analysis_id]
    assert result.report_ids == []
    assert not list((store.root / 'report').rglob('report.*'))


def test_no_analysis_reason_and_bounded_default_budget(context, monkeypatch):
    store, task, _, _ = context
    result = ReportService(store).generate(task, [])
    assert result.status == 'need_analysis'
    assert result.validation_summary['reason_code'] == 'missing_analysis'
    monkeypatch.delenv('MAX_TOOL_CALLS', raising=False)
    defaults = Settings(_env_file=None)
    assert defaults.max_tool_calls == 256
    assert defaults.task_timeout_seconds == 3600
    assert defaults.max_agent_steps == 1024
    assert defaults.sandbox_timeout_seconds == 120
