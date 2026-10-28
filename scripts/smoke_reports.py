"""Opt-in API -> persistent Worker -> Agents -> two versioned report deliveries.

No model calls without --live. Requires services already started by start.ps1.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import datetime
import hashlib
import io
import json
from pathlib import Path
import sys
import time
from uuid import uuid4
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def expected_sources(root):
    expected = {}
    for name in ('winequality-red.csv', 'winequality-white.csv'):
        raw = (root / 'KB' / name).read_bytes()
        rows = list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig')), delimiter=';'))
        quality = [int(row['quality']) for row in rows]
        expected[name] = dict(hash=digest(raw), rows=len(rows), quality_mean=sum(quality) / len(quality),
                              quality_distribution=dict(Counter(quality)))
    return expected


def poll(client, task_id, timeout):
    deadline = time.monotonic() + timeout
    last = None
    while time.monotonic() < deadline:
        task = client.get(f'tasks/{task_id}').json()
        status = task['status']
        if status != last:
            print(f'Task {task_id}: {status}', flush=True)
            last = status
        if status in {'succeeded', 'failed', 'cancelled', 'interrupted', 'waiting_user'}:
            require(status == 'succeeded', f'Task {task_id} ended in {status}; inspect local task record.')
            return task
        time.sleep(2)
    # Leave the task record intact for diagnosis, but stop unexpected spend.
    client.post(f'tasks/{task_id}/cancel')
    raise TimeoutError(f'Task {task_id} exceeded smoke timeout and cancellation was requested.')


def verify_delivery(client, task, expected, directory):
    from bs4 import BeautifulSoup
    from openpyxl import load_workbook
    from src.chart_svg import chart_csp, inline_svg, interaction_script

    task_id = task['task_id']
    artifacts = client.get(f'tasks/{task_id}/artifacts').json()['items']
    by_id = {a['artifact_id']: a for a in artifacts}
    directory.mkdir(parents=True, exist_ok=True)
    cache = {}

    def download(artifact):
        aid = artifact['artifact_id']
        if aid not in cache:
            content = client.get(f'artifacts/{aid}/download').content
            require(digest(content) == artifact['hash'], 'HTTP artifact SHA256 mismatch')
            cache[aid] = content
        return cache[aid]

    manifests = [a for a in artifacts if a['kind'] == 'report_manifest']
    require(len(manifests) == 1, 'Expected exactly one registered report manifest')
    manifest_bytes = download(manifests[0])
    manifest = json.loads(manifest_bytes)
    require(manifest['validation']['status'] == 'passed', 'Report validation did not pass')
    require(set(manifest['formats']) == {'html', 'xlsx'}, 'Both report formats are required')
    require(manifest['task_id'] == task_id, 'Manifest task mismatch')
    require(set(a['analysis_id'] for a in manifest['analyses']) == set(task['result']['analysis_ids']), 'Result and manifest analyses differ')
    (directory / 'manifest.json').write_bytes(manifest_bytes)
    sources = {source['file_id']: source for source in manifest['sources']}
    require(set(source['name'] for source in sources.values()) == set(expected), 'Wrong report input files')
    for source in sources.values():
        require(source['hash'] == expected[source['name']]['hash'], 'Report source differs from independent CSV')

    for artifact in manifest['source_artifacts']:
        download(artifact)
    numeric_verified, distributions_verified, chart_ids = set(), set(), set()
    for analysis in manifest['analyses']:
        names = [sources[fid]['name'] for fid in analysis['input_versions']]
        if len(names) != 1:
            continue
        name = names[0]
        metrics = analysis['metrics']
        if 'quality_mean' in metrics:
            require(metrics.get('source_rows') == expected[name]['rows'], 'Full source row count mismatch')
            require(abs(metrics['quality_mean'] - expected[name]['quality_mean']) < 1e-9, 'Independent quality mean mismatch')
            numeric_verified.add(name)
        for aid in analysis['table_ids']:
            rows = list(csv.DictReader(io.StringIO(cache[aid].decode('utf-8-sig'))))
            if rows and 'quality' in rows[0] and 'quality_count' in rows[0]:
                counts = {int(float(row['quality'])): int(float(row['quality_count'])) for row in rows}
                require(counts == expected[name]['quality_distribution'], 'Independent quality distribution mismatch')
                require(bool(analysis['chart_ids']), 'Distribution has no registered chart')
                distributions_verified.add(name)
        for aid in analysis['chart_ids']:
            chart_ids.add(aid)
    require(numeric_verified == set(expected), 'Both full-data means and row counts are required')
    require(distributions_verified == set(expected), 'Both quality distributions and charts are required')

    reports = [a for a in artifacts if a['kind'] == 'report']
    require(len(reports) == 2, 'Expected two report artifacts')
    report_bytes = {Path(a['name']).suffix: download(a) for a in reports}
    require(set(report_bytes) == {'.html', '.xlsx'}, 'Report artifact suffixes differ')
    for entry in manifest['files']:
        content = report_bytes[Path(entry['name']).suffix]
        require(digest(content) == entry['sha256'], 'Manifest report hash mismatch')
        (directory / entry['name']).write_bytes(content)
    bindings = {(m['analysis_id'], m['key']): m['value'] for m in manifest['metric_bindings']}
    # Browsers and the application's text reader normalize Windows CRLF to LF.
    soup = BeautifulSoup(report_bytes['.html'].decode('utf-8').replace('\r\n', '\n'), 'html.parser')
    actual = {(cell['data-analysis-id'], cell['data-metric-key']): cell.get_text() for cell in soup.select('[data-metric-key]')}
    show_metrics = manifest['presentation']['layout']['include_metrics_summary']
    require(actual == (bindings if show_metrics else {}), 'HTML metric summary differs from configured layout')
    require(not show_metrics, 'New reports must hide the metric summary by default')
    scripts = soup.find_all('script')
    require(len(scripts) == (1 if chart_ids else 0), 'Unexpected chart script count')
    require(all(s.get('id') == 'ds-chart-interactions' and s.get_text() == interaction_script()
                and set(s.attrs) == {'id'} for s in scripts), 'Untrusted chart script')
    require(soup.find('meta', attrs={'http-equiv': 'Content-Security-Policy'})['content'] == chart_csp(), 'HTML CSP mismatch')
    require(len(soup.select('svg[data-ds-chart]')) == len(chart_ids), 'HTML vector chart count mismatch')
    source_artifacts = {a['artifact_id']: a for a in manifest['source_artifacts']}
    chart_hashes = set()
    for cid in chart_ids:
        source = source_artifacts[cid]
        companions = [a for a in source_artifacts.values() if a['source_ids'] == [cid]]
        svg = source if source['name'].endswith('.svg') else next(a for a in companions if a['kind'] == 'chart_svg')
        png = source if source['name'].endswith('.png') else next(a for a in companions if a['kind'] == 'chart_preview')
        expected_svg = BeautifulSoup(str(inline_svg(cache[svg['artifact_id']], cid)), 'html.parser').find('svg')
        actual_svg = soup.find('svg', attrs={'data-ds-chart': cid})
        require(str(actual_svg) == str(expected_svg), 'HTML SVG differs from registered vector source')
        chart_hashes.add(digest(cache[png['artifact_id']]))
    book = load_workbook(io.BytesIO(report_bytes['.xlsx']), data_only=False)
    try:
        require(('指标汇总' in book.sheetnames) == show_metrics, 'Excel metric summary differs from configured layout')
        if show_metrics:
            excel_bindings = {(row[0], row[1]): row[3] for row in list(book['指标汇总'].values)[1:]}
            require(excel_bindings == bindings, 'Excel values differ from manifest bindings')
        require(all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row), 'Excel contains formulas')
        for index, table in enumerate(manifest['tables'], 1):
            actual_rows = sum(book[name].max_row - 1 for name in book.sheetnames if name.startswith(f'结果表_{index}_'))
            require(actual_rows == table['total_rows'], 'Excel omitted result rows')
    finally:
        book.close()
    with zipfile.ZipFile(io.BytesIO(report_bytes['.xlsx'])) as archive:
        embedded = {digest(archive.read(name)) for name in archive.namelist() if name.startswith('xl/media/')}
    require(chart_hashes.issubset(embedded), 'Excel lacks source chart bytes')
    return dict(task_id=task_id, version=manifest['report_version'], analysis_ids=task['result']['analysis_ids'],
                report_artifacts=[{k: a[k] for k in ('artifact_id', 'relative_path', 'hash')} for a in reports],
                metrics={name: {key: expected[name][key] for key in ('rows', 'quality_mean')} for name in expected},
                tool_calls=task['result'].get('tool_calls'), verified=True), artifacts


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--live', action='store_true', help='Authorize paid models via the running API and Worker.')
    parser.add_argument('--api-url', default=None)
    parser.add_argument('--resume-task', help='Verify an existing first report task before submitting only the reuse task.')
    parser.add_argument('--timeout', type=int, default=1020, help='Maximum seconds for each task, then request cancellation.')
    args = parser.parse_args()
    if not args.live:
        print('No model calls made. Start API/Worker, then pass --live for the paid report acceptance test.')
        return
    import httpx
    from src.config import Settings
    settings = Settings()
    expected = expected_sources(settings.root)
    directory = ROOT / '.runtime' / 'smoke_reports' / (datetime.now().strftime('%Y%m%dT%H%M%S') + '_' + uuid4().hex[:8])
    directory.mkdir(parents=True)
    summary = {'status': 'running', 'tool_budget': settings.max_tool_calls, 'deliveries': []}

    def save():
        (directory / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')

    try:
        with httpx.Client(base_url=(args.api_url or settings.api_url).rstrip('/') + '/api/v1/',
                          timeout=30, trust_env=False, event_hooks={'response': [lambda response: response.raise_for_status()]}) as client:
            health = client.get('health').json()
            require(health['worker']['ready'] and health['models_configured'], 'API Worker/model configuration is not ready')
            if args.resume_task:
                original = client.get(f'tasks/{args.resume_task}').json()
                sid = original['session_id']
            else:
                sid = client.post('sessions', json={'title': 'Wine report API acceptance'}).json()['session_id']
            summary['session_id'] = sid
            files = client.get(f'sessions/{sid}/files').json()['items']
            chosen = [file for file in files if file['source'] == 'kb' and file['name'] in expected]
            require(len(chosen) == 2, 'Both wine datasets must be registered in the running workspace KB')
            for file in chosen:
                require(file['hash'] == expected[file['name']]['hash'], 'API KB does not match local independent input')
            file_ids = [file['file_id'] for file in chosen]
            request = ('请比较选中的红、白葡萄酒：分别给出完整数据行数、quality平均值，以及quality各评分的数量分布柱状图。'
                       '只做这些内容，不增加其他统计。请一次委派数据分析专家完成两份文件的全部分析，'
                       '每份文件保留均值结果和带柱状图的评分分布结果，共四份分析结果。'
                       '然后由报告专家将这些已有结果、两张图及来源整理为HTML和Excel报告。'
                       '请复用已得到的分析，不要重复计算。')
            first_id = args.resume_task or client.post(f'sessions/{sid}/messages', json={'content': request, 'file_ids': file_ids,
                'expected_outputs': ['html', 'xlsx']}, headers={'Idempotency-Key': uuid4().hex}).json()['task_id']
            summary['first_task_id'] = first_id
            save()
            first_task = poll(client, first_id, args.timeout)
            first, _ = verify_delivery(client, first_task, expected, directory / 'first')
            require(first['tool_calls'] <= settings.max_tool_calls, 'Configured tool budget exceeded')
            summary['deliveries'].append(first)
            save()
            second_id = client.post(f'sessions/{sid}/messages', json={'content': '仅使用指定已有analysis_ids重新生成管理层版HTML和XLSX报告。'
                '复用全部已有指标、结果表和图表，不要重新分析数据，不要委派analyst。请报告专家使用不同标题并保留来源。',
                'file_ids': file_ids, 'analysis_ids': first['analysis_ids'], 'expected_outputs': ['html', 'xlsx']},
                headers={'Idempotency-Key': uuid4().hex}).json()['task_id']
            summary['second_task_id'] = second_id
            save()
            second_task = poll(client, second_id, args.timeout)
            second, artifacts = verify_delivery(client, second_task, expected, directory / 'second')
            require(second['tool_calls'] <= settings.max_tool_calls, 'Configured tool budget exceeded')
            require(set(second['analysis_ids']) == set(first['analysis_ids']), 'Report-only task did not reuse exact analyses')
            require(not any(a['kind'] in {'table', 'chart', 'script'} for a in artifacts), 'Report-only task created new analysis artifacts')
            events = client.get(f'tasks/{second_id}/events').json()['items']
            require(not any(e.get('role') == 'analyst' for e in events), 'Report-only task invoked analyst')
            require(first['version'] != second['version'], 'Report version was overwritten')
            first_paths = {a['relative_path'] for a in first['report_artifacts']}
            require(first_paths.isdisjoint(a['relative_path'] for a in second['report_artifacts']), 'Report paths were reused')
            for artifact in first['report_artifacts']:
                require(digest(client.get(f"artifacts/{artifact['artifact_id']}/download").content) == artifact['hash'], 'Earlier report changed after new version')
            summary['deliveries'].append(second)
            summary['status'] = 'passed'
            save()
            print(json.dumps(summary, ensure_ascii=False, indent=2))
            print(f'PASS: API/Worker, independent CSV/manifest/HTML/XLSX/download verification; evidence: {directory}')
    except Exception as exc:
        summary['status'] = 'failed'
        summary['error_type'] = type(exc).__name__
        if isinstance(exc, (AssertionError, TimeoutError)):
            summary['check'] = str(exc)
        save()
        # Provider exception bodies and model text are deliberately excluded.
        print(f"FAIL: {summary.get('check', type(exc).__name__)}; safe evidence saved to {directory}")
        raise SystemExit(1) from None


if __name__ == '__main__':
    main()
